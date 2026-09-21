from __future__ import annotations

import io
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation

from backend.exports import build_output

ROOT = Path(__file__).resolve().parents[1]
TEMPLATES = ROOT / "data" / "REBUILD_INPUT_v1" / "REBUILD_INPUT_v1" / "03_OUTPUT_TEMPLATES"


def _source(path: str = "02_ITB.pdf", locator: str = "page 2", quote: str = "450 m3/day") -> dict:
    return {"source_id": "doc-1:block-1", "document_id": "doc-1", "source_path": path, "sha256": "a" * 64, "revision": "Rev01", "approval_status": "APPROVED", "locator_type": "pdf_page", "locator": locator, "quote": quote, "url": "http://127.0.0.1:8780/#source=doc-1"}


def test_itb_export_preserves_template_and_adds_key_facts_with_sources() -> None:
    draft = {
        "kind": "itb",
        "title": "Blue Line ITB analysis",
        "items": [{"item_id": "ITB-01", "itb_clause": "Discharge", "current_condition": "450 m3/day", "historical_case": "Fresh-water settlement only", "differences": "Saline site", "decision": "ADAPT", "source_refs": [_source()], "review_required": "Confirm treatment design"}],
        "key_facts": [{"fact": "Discharge allowance", "value": 450, "unit": "m3/day", "source_refs": [_source()]}],
    }
    payload, mime, extension = build_output(draft, TEMPLATES / "ITB_Analysis_Template.xlsx")
    workbook = load_workbook(io.BytesIO(payload), data_only=False)
    sheet = workbook["ITB"]
    assert mime == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert extension == ".xlsx"
    assert [sheet.cell(5, column).value for column in range(1, 9)] == ["항목 ID", "ITB 조항", "신규 조건", "과거 사례", "조건 차이", "적용 판단", "원문 근거", "확인 필요 사항"]
    assert sheet["A6"].value == "ITB-01"
    assert "02_ITB.pdf" in sheet["G6"].value and "page 2" in sheet["G6"].value
    facts = workbook["Key Facts"]
    assert facts["A2"].value == "Discharge allowance"
    assert facts["B2"].value == 450 and facts["C2"].value == "m3/day"
    assert "Sources" in workbook.sheetnames
    sources = workbook["Sources"]
    assert sheet["G6"].hyperlink.target == "#'Sources'!A2"
    assert sources["J2"].hyperlink.target == "http://127.0.0.1:8780/#source=doc-1"
    assert sources["D2"].value == "a" * 64


def test_itb_empty_evidence_is_review_required() -> None:
    payload, _, _ = build_output({"type": "itb", "items": [{"item_id": "ITB-EMPTY", "itb_clause": "Unknown"}]}, TEMPLATES / "ITB_Analysis_Template.xlsx")
    sheet = load_workbook(io.BytesIO(payload))["ITB"]
    assert sheet["A6"].value == "ITB-EMPTY"
    assert sheet["G6"].value == "REVIEW_REQUIRED"
    assert "REVIEW_REQUIRED" in sheet["H6"].value


def test_risk_export_leaves_unknown_rating_fields_blank_and_reopens() -> None:
    draft = {"kind": "risk", "items": [{"risk_id": "R-01", "risk": "Saline discharge treatment", "cause": "Site conditions differ", "impact": "Permit delay", "mitigation": "Confirm treatment train", "source_refs": [_source("09_Field.txt", "line 4", "saline groundwater")]}]}
    payload, mime, extension = build_output(draft, TEMPLATES / "Risk_Register_Template.xlsx")
    workbook = load_workbook(io.BytesIO(payload), data_only=False)
    sheet = workbook.active
    values = [[cell.value for cell in row] for row in sheet.iter_rows()]
    flat = [value for row in values for value in row if value is not None]
    assert mime.endswith("spreadsheetml.sheet") and extension == ".xlsx"
    assert "R-01" in flat and any("REVIEW_REQUIRED" in str(value) for value in flat)
    risk_row = next(cell.row for cell in sheet["A"] if cell.value == "R-01")
    assert sheet.cell(risk_row, 9).hyperlink.target == "#'Sources'!A2"
    row = next(row for row in values if "R-01" in row)
    header_row = next(row for row in values if any("발생" in str(v) or "가능" in str(v) for v in row if v))
    for index, header in enumerate(header_row):
        if header and any(token in str(header) for token in ("발생가능성", "가능성", "강도", "점수", "Probability", "Score")):
            assert row[index] in (None, "")


def test_review_deck_uses_five_slide_template_and_embeds_sources_and_unknowns() -> None:
    draft = {
        "kind": "slides", "title": "Blue Line review",
        "business": [{"text": "Package 4 overview", "source_refs": [_source("brief.docx", "paragraph 2", "Package 4")]}],
        "itb": [{"text": "Discharge allowance 450 m3/day", "source_refs": [_source()]}],
        "applicability": [{"text": "Adapt settlement method for salinity", "source_refs": [_source("past.docx", "paragraph 15", "fresh water")]}],
        "risks": [{"text": "Treatment design unconfirmed", "source_refs": [_source("field.txt", "line 4", "saline")]}],
        "unknowns": [{"text": "Pump capacity", "source_refs": []}],
    }
    payload, mime, extension = build_output(draft, TEMPLATES / "Review_Deck_Template.pptx")
    deck = Presentation(io.BytesIO(payload))
    all_text = "\n".join(shape.text for slide in deck.slides for shape in slide.shapes if hasattr(shape, "text_frame"))
    assert mime == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    assert extension == ".pptx" and len(deck.slides) == 5
    assert "Package 4 overview" in all_text and "450 m3/day" in all_text
    assert "02_ITB.pdf" in all_text and "REVIEW_REQUIRED" in all_text
    links = [run.hyperlink.address for slide in deck.slides for shape in slide.shapes if hasattr(shape, "text_frame") for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if run.hyperlink.address]
    assert "http://127.0.0.1:8780/#source=doc-1" in links


def _comparison_row() -> dict:
    return {
        "id": "discharge",
        "title": "방류 조건",
        "current": "허용량 450 m3/day",
        "past": "과거 담수 침전 사례",
        "differences": ["염수 조건", "허용량과 설계량 구분"],
        "decision": "REVIEW_REQUIRED",
        "rationale": "동일 의미인지 설계 검토가 필요합니다.",
        "missing_information": ["펌프 설계량 미확정"],
        "source_refs": [_source()],
        "severity": None,
        "mitigation": "처리 공정을 확인하세요.",
    }


def test_server_comparison_rows_map_to_itb_columns_without_serializing_lists() -> None:
    payload, _, _ = build_output({"kind": "itb", "rows": [_comparison_row()]}, TEMPLATES / "ITB_Analysis_Template.xlsx")
    sheet = load_workbook(io.BytesIO(payload))["ITB"]
    assert sheet["A6"].value == "discharge"
    assert sheet["B6"].value == "방류 조건"
    assert sheet["C6"].value == "허용량 450 m3/day"
    assert sheet["D6"].value == "과거 담수 침전 사례"
    assert "염수 조건" in sheet["E6"].value
    assert "펌프 설계량 미확정" in sheet["H6"].value


def test_server_comparison_rows_populate_all_review_deck_sections() -> None:
    payload, _, _ = build_output({"kind": "slides", "title": "심의 초안", "rows": [_comparison_row()]}, TEMPLATES / "Review_Deck_Template.pptx")
    deck = Presentation(io.BytesIO(payload))
    slide_text = ["\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text_frame")) for slide in deck.slides]
    assert "방류 조건" in slide_text[0]
    assert "허용량 450 m3/day" in slide_text[1]
    assert "염수 조건" in slide_text[2]
    assert "처리 공정을 확인하세요." in slide_text[3]
    assert "펌프 설계량 미확정" in slide_text[4]

def test_user_edited_review_fields_survive_all_output_formats() -> None:
    row = _comparison_row()
    row.update({
        "decision": "ADAPT_AFTER_REVIEW",
        "rationale": "Browser reviewer verified the uploaded notice period.",
        "mitigation": "Preserve the reviewer mitigation exactly.",
        "missing_information": ["Confirm the final permit owner."],
    })

    itb_bytes, _, _ = build_output({"kind": "itb", "rows": [row]}, TEMPLATES / "ITB_Analysis_Template.xlsx")
    itb = load_workbook(io.BytesIO(itb_bytes))["ITB"]
    assert "ADAPT_AFTER_REVIEW" in itb["F6"].value
    assert "Browser reviewer verified the uploaded notice period." in itb["F6"].value
    assert "Confirm the final permit owner." in itb["H6"].value

    risk_bytes, _, _ = build_output({"kind": "risk", "rows": [row]}, TEMPLATES / "Risk_Register_Template.xlsx")
    risk = load_workbook(io.BytesIO(risk_bytes)).active
    assert "Browser reviewer verified the uploaded notice period." in risk["C6"].value
    assert risk["G6"].value == "Preserve the reviewer mitigation exactly."
    assert "Confirm the final permit owner." in risk["J6"].value

    slide_bytes, _, _ = build_output({"kind": "slides", "rows": [row]}, TEMPLATES / "Review_Deck_Template.pptx")
    deck = Presentation(io.BytesIO(slide_bytes))
    text = "\n".join(shape.text for slide in deck.slides for shape in slide.shapes if hasattr(shape, "text_frame"))
    assert "Browser reviewer verified the uploaded notice period." in text
    assert "Preserve the reviewer mitigation exactly." in text
    assert "Confirm the final permit owner." in text