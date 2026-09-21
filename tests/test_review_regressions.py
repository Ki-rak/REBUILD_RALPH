from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pptx import Presentation

from backend.config import Settings
from backend.exports import build_output
from backend.server import create_app
from test_api import flow, project, upload, draft, auth
from test_exports import TEMPLATES


def test_authentication_client_closed_when_session_rejected(monkeypatch):
    clients = []

    class RejectedAuth:
        def __init__(self, *args):
            self.closed = False
            clients.append(self)
        def get_user(self, token):
            raise RuntimeError("credential-bearing upstream diagnostic")
        def close(self):
            self.closed = True

    monkeypatch.setattr("backend.server.SupabaseAuth", RejectedAuth)
    settings = Settings("https://wsziosnttnxefgfbgpeq.supabase.co", "test-key",
                        "local", "test-integrity-key-of-at-least-32-characters")
    response = TestClient(create_app(settings), raise_server_exceptions=False).get(
        "/api/projects", headers=auth())
    assert response.status_code == 401
    assert "credential-bearing" not in response.text
    assert clients and all(client.closed for client in clients)


def test_risk_without_calculation_metadata_explicitly_requests_recalculation(tmp_path):
    workbook = load_workbook(TEMPLATES / "Risk_Register_Template.xlsx")
    workbook.calculation = None
    template = tmp_path / "Risk_Register_Template.xlsx"
    workbook.save(template)
    data, _, _ = build_output({"kind": "risk", "rows": []}, template)
    result = load_workbook(BytesIO(data))
    assert result.calculation is not None
    assert result.calculation.fullCalcOnLoad and result.calculation.forceFullCalc
    assert result.calculation.calcMode == "auto"


def test_five_slide_draft_edits_render_on_corresponding_slide_and_invalidate(flow):
    client, _, _ = flow
    past = project(client, "Historical", "historical")
    upload(client, past, "Notice period: 28 calendar days.")
    current = project(client)
    upload(client, current, "Employer: Blue Transit.\nNotice period: 17 calendar days.")
    value = draft(client, current, "slides")
    expected = ["overview", "itb", "applicability", "risks", "decisions"]
    assert [row["id"] for row in value["rows"]] == expected
    for index, row in enumerate(value["rows"]):
        row["title"] = f"Reviewed heading {index+1}"
        row["current"] = f"Reviewed content for slide {index+1}"
        row["rationale"] = f"Review rationale {index+1}"
    changed = client.patch(f"/api/drafts/{value['id']}",
        json={"rows": value["rows"], "revision": value["revision"]}, headers=auth())
    assert changed.status_code == 200
    value = changed.json()
    approved = client.post(f"/api/drafts/{value['id']}/approve",
        json={"confirmed": True, "revision": value["revision"]}, headers=auth())
    assert approved.status_code == 200
    response = client.post(f"/api/drafts/{value['id']}/export", headers=auth())
    assert response.status_code == 200
    deck = Presentation(BytesIO(response.content))
    assert len(deck.slides) == 5
    for index, slide in enumerate(deck.slides):
        text = "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text"))
        assert f"Reviewed heading {index+1}" in text
        assert f"Reviewed content for slide {index+1}" in text
        assert f"Review rationale {index+1}" in text
    value["rows"][0]["current"] = "Second review"
    assert client.patch(f"/api/drafts/{value['id']}",
        json={"rows": value["rows"], "revision": value["revision"]},
        headers=auth()).status_code == 200
    assert client.post(f"/api/drafts/{value['id']}/export", headers=auth()).status_code == 409

def test_excel_source_and_review_strings_remain_literal_not_executable_formula():
    formula_text = '=HYPERLINK("https://example.invalid","literal source")'
    reference = {"source_id": "doc:b", "document_id": "doc", "source_path": formula_text,
                 "sha256": "a"*64, "locator": "Sheet1!A1", "quote": formula_text}
    value = {"id": "item", "title": formula_text, "current": formula_text,
             "source_refs": [reference], "severity": 3, "probability": 2}
    for kind, filename in [("itb", "ITB_Analysis_Template.xlsx"),
                           ("risk", "Risk_Register_Template.xlsx")]:
        data, _, _ = build_output({"kind": kind, "rows": [value]}, TEMPLATES / filename)
        workbook = load_workbook(BytesIO(data), data_only=False)
        matches = [cell for sheet in workbook for row in sheet for cell in row
                   if cell.value == formula_text]
        assert matches and all(cell.data_type == "s" for cell in matches)
        if kind == "risk":
            assert workbook.active["F6"].data_type == "f"
            assert workbook.active["F6"].value.endswith("D6*E6))")

def test_draft_list_uses_same_public_output_kind_as_detail(flow):
    client, _, _ = flow
    pid = project(client)
    upload(client, pid, "Notice period: 19 calendar days.")
    value = draft(client, pid, "itb")
    listed = client.get(f"/api/projects/{pid}/drafts", headers=auth())
    assert listed.status_code == 200
    saved = next(item for item in listed.json() if item["id"] == value["id"])
    assert saved["kind"] == value["kind"] == "itb"
    assert client.get(f"/api/drafts/{value['id']}", headers=auth()).json()["kind"] == "itb"

def test_concurrent_upload_during_analysis_cannot_approve_stale_rows(flow, monkeypatch):
    import backend.server as server
    client, _, _ = flow
    pid = project(client)
    upload(client, pid, "Notice period: 14 calendar days.", "before.txt")
    original_compare = server.compare
    injected = [False]
    def changed_while_analyzing(current, historical):
        result = original_compare(current, historical)
        if not injected[0]:
            injected[0] = True
            upload(client, pid, "Employer: Newly arrived during analysis.", "during.txt")
        return result
    monkeypatch.setattr(server, "compare", changed_while_analyzing)
    response = client.post(f"/api/projects/{pid}/drafts",
        json={"kind": "itb", "mode": "rules"}, headers=auth())
    assert response.status_code == 409
    assert response.json()["detail"]["code"] == "INPUT_CHANGED"
    assert client.get(f"/api/projects/{pid}/drafts", headers=auth()).json() == []