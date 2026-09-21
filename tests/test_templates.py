from io import BytesIO
from hashlib import sha256

from openpyxl import load_workbook

from test_api import flow, project, upload, draft, auth
from test_exports import TEMPLATES


def test_template_version_is_reusable_preserved_and_requires_reapproval(flow):
    client, records, files = flow
    source = (TEMPLATES / "ITB_Analysis_Template.xlsx").read_bytes()
    original_hash = sha256(source).hexdigest()
    workbook = load_workbook(BytesIO(source))
    workbook["ITB"]["A1"] = "Registered custom version"
    output = BytesIO()
    workbook.save(output)
    changed = output.getvalue()
    response = client.post("/api/templates/itb/versions",
        files={"file": ("ITB_custom.xlsx", changed,
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=auth())
    assert response.status_code == 201, response.text
    version = response.json()["version"]
    assert version == sha256(changed).hexdigest() and version != original_hash
    templates = client.get("/api/templates", headers=auth()).json()
    assert len(templates) == 7
    versions = next(item for item in templates if item["id"] == "itb")["versions"]
    assert {item["version"] for item in versions} >= {version, original_hash}
    assert client.get(f"/api/templates/itb/original?version={original_hash}", headers=auth()).content == source
    assert client.get(f"/api/templates/itb/original?version={version}", headers=auth()).content == changed
    assert client.get(f"/api/templates/itb/original?version={version}", headers=auth("bob")).status_code == 404
    for index in range(2):
        pid = project(client, f"Reuse {index}")
        upload(client, pid, "Notice period: 17 calendar days.")
        value = draft(client, pid)
        assert client.post(f"/api/drafts/{value['id']}/approve",
            json={"confirmed": True, "revision": 1}, headers=auth()).status_code == 200
        edit = client.patch(f"/api/drafts/{value['id']}",
            json={"rows": value["rows"], "revision": 1, "template_version": version},
            headers=auth())
        assert edit.status_code == 200
        value = edit.json()
        assert value["status"] == "draft" and value["template_sha256"] == version
        assert client.post(f"/api/drafts/{value['id']}/export", headers=auth()).status_code == 409
        assert client.post(f"/api/drafts/{value['id']}/approve",
            json={"confirmed": True, "revision": value["revision"]}, headers=auth()).status_code == 200
        rendered = client.post(f"/api/drafts/{value['id']}/export", headers=auth())
        assert rendered.status_code == 200
        assert load_workbook(BytesIO(rendered.content))["ITB"]["A1"].value == "Registered custom version"
    assert (TEMPLATES / "ITB_Analysis_Template.xlsx").read_bytes() == source


def test_invalid_template_does_not_replace_existing_versions(flow):
    client, _, _ = flow
    result = client.post("/api/templates/itb/versions",
        files={"file": ("broken.xlsx", b"corrupt", "application/octet-stream")}, headers=auth())
    assert result.status_code == 422
    assert len(client.get("/api/templates", headers=auth()).json()) == 7

def test_template_upload_retry_recovers_after_storage_succeeded_database_failed(flow, monkeypatch):
    from test_api import FakeStorage
    client, _, _ = flow
    save = FakeStorage.save
    failed = [False]
    def fail_once(self, kind, *args, **kwargs):
        if kind == "template" and not failed[0]:
            failed[0] = True
            raise RuntimeError("synthetic transient write failure")
        return save(self, kind, *args, **kwargs)
    monkeypatch.setattr(FakeStorage, "save", fail_once)
    data = (TEMPLATES / "ITB_Analysis_Template.xlsx").read_bytes()
    send = lambda: client.post("/api/templates/itb/versions",
        files={"file": ("retry.xlsx", data, "application/octet-stream")}, headers=auth())
    assert send().status_code == 503
    assert send().status_code == 201


def test_forged_template_version_metadata_is_rejected(flow):
    client, records, _ = flow
    data = (TEMPLATES / "ITB_Analysis_Template.xlsx").read_bytes()
    response = client.post("/api/templates/itb/versions",
        files={"file": ("registered.xlsx", data, "application/octet-stream")}, headers=auth())
    assert response.status_code == 201
    key = next(key for key in records if key[1].startswith("template:"))
    records[key]["payload"]["filename"] = "forged.xlsx"
    denied = client.get(f"/api/templates/itb/original?version={sha256(data).hexdigest()}", headers=auth())
    assert denied.status_code == 409

def test_slide_template_rejects_changed_canvas_without_persisting(flow):
    from pptx import Presentation
    from pptx.util import Inches
    client, records, files = flow
    original = (TEMPLATES / "Review_Deck_Template.pptx").read_bytes()
    deck = Presentation(BytesIO(original))
    deck.slide_width, deck.slide_height = Inches(7.5), Inches(10)
    candidate = BytesIO()
    deck.save(candidate)
    records_before, files_before = len(records), len(files)
    response = client.post("/api/templates/slides/versions",
        files={"file": ("portrait.pptx", candidate.getvalue(), "application/octet-stream")},
        headers=auth())
    assert response.status_code == 422
    assert response.json()["detail"]["code"] == "SLIDE_TEMPLATE_SIZE_UNSUPPORTED"
    assert "크기" in response.json()["detail"]["message"]
    assert len(records) == records_before and len(files) == files_before
    assert (TEMPLATES / "Review_Deck_Template.pptx").read_bytes() == original


def test_slide_template_same_canvas_custom_metadata_remains_accepted(flow):
    from pptx import Presentation
    client, _, _ = flow
    deck = Presentation(TEMPLATES / "Review_Deck_Template.pptx")
    deck.core_properties.title = "Reusable committee template"
    candidate = BytesIO()
    deck.save(candidate)
    response = client.post("/api/templates/slides/versions",
        files={"file": ("committee-custom.pptx", candidate.getvalue(), "application/octet-stream")},
        headers=auth())
    assert response.status_code == 201
    version = response.json()["version"]
    assert client.get(f"/api/templates/slides/original?version={version}", headers=auth()).content == candidate.getvalue()
