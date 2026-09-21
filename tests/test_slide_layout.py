from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

import pytest
from openpyxl import load_workbook
from pptx import Presentation
from pptx.util import Inches, Pt

import backend.drafts as slide_contract
from backend.analysis import compare
from backend.drafts import build_slide_rows, validate_slide_rows
from backend.exports import build_output
from test_analysis import actual_p01_n01
from test_exports import TEMPLATES


def source(index: int) -> dict:
    return {
        "source_id": f"doc-{index}:block-{index}",
        "document_id": f"doc-{index}",
        "filename": f"evidence-{index}.pdf",
        "source_path": f"folder/evidence-{index}.pdf",
        "sha256": str(index) * 64,
        "revision": "Rev01",
        "approval_status": "APPROVED",
        "locator_type": "pdf_page",
        "locator": f"page {index}",
        "quote": f"Complete evidence quote {index}",
        "url": f"http://127.0.0.1:8780/#source=doc-{index}",
    }


def test_wrapped_layout_budget_rejects_long_korean_title_and_body():
    rows = build_slide_rows([])
    rows[0]["title"] = "긴" * 80
    with pytest.raises(ValueError, match="SLIDE_TITLE_TOO_LONG"):
        validate_slide_rows(rows)

    rows = build_slide_rows([])
    rows[0]["current"] = "\n".join(["가" * 50] * 12)
    assert len(rows[0]["current"]) <= 650
    with pytest.raises(ValueError, match="SLIDE_CONTENT_TOO_LONG"):
        validate_slide_rows(rows)

    rows = build_slide_rows([])
    rows[0].update({"title": "W" * 80, "current": "\n".join(["W" * 88] * 7),
                    "decision": "", "rationale": "", "missing_information": []})
    with pytest.raises(ValueError, match="SLIDE_TITLE_TOO_LONG"):
        validate_slide_rows(rows)

    rows = build_slide_rows([])
    rows[0].update({"title": "fit", "current": "\n".join(["H" * 88] * 7),
                    "decision": "", "rationale": "", "missing_information": []})
    with pytest.raises(ValueError, match="SLIDE_CONTENT_TOO_LONG"):
        validate_slide_rows(rows)


def test_nested_severity_samples_are_flattened_with_roles_and_statistics():
    risk_ref = source(71)
    severity_ref = source(72)
    comparison = {
        "id": "discharge", "title": "Discharge", "current": "180 mg/L",
        "rationale": "generic rationale " * 30, "mitigation": "Add treatment stage", "severity": 4,
        "value_classes": {"current": [{"kind": "new_design"}], "historical": [{"kind": "historical_observation"}]},
        "source_refs": [], "current_refs": [], "historical_refs": [],
        "severity_evidence": {
            "sample_count": 1, "project_count": 1, "unique_case_count": 1,
            "distribution": {"4": 1}, "scale": "1-5", "reason": "sample",
            "samples": [{"project_id": "past-71", "severity": 4, "scale": "1-5",
                         "value_classes": ["observed"], "risk_refs": [risk_ref],
                         "severity_refs": [severity_ref]}],
        },
    }
    rows = build_slide_rows([comparison])
    keys = {(ref.get("source_id"), ref.get("locator")) for ref in rows[1]["source_refs"]}
    assert (risk_ref["source_id"], risk_ref["locator"]) in keys
    assert (severity_ref["source_id"], severity_ref["locator"]) in keys
    detail = rows[1]["detail_text"]
    assert "past-71" in detail and "observed" in detail
    assert "new_design" in detail and "historical_observation" in detail
    assert "risk_refs" in detail and "severity_refs" in detail
    risk_visible = rows[3]["current"]
    assert "180 mg/L" in risk_visible and "Add treatment stage" in risk_visible and "강도 4" in risk_visible
    assert "generic rationale" not in risk_visible and "generic rationale" in rows[3]["detail_text"]


def test_real_p01_n01_deck_keeps_full_detail_and_refs_in_editable_notes():
    current, historical = actual_p01_n01()
    comparisons = compare(current, historical)["rows"]
    rows = build_slide_rows(comparisons)
    validate_slide_rows(rows)
    assert all(isinstance(row["detail_text"], str) and row["detail_text"] for row in rows)
    assert all(row["excerpted"] is True for row in rows)
    expected_extra_refs = [ref for comparison in comparisons
                           for field in ("excluded_refs", "mitigation_refs")
                           for ref in comparison.get(field, [])]
    normalized_keys = {(ref.get("source_id"), ref.get("locator"), ref.get("quote"))
                       for row in rows for ref in row["source_refs"]}
    assert expected_extra_refs
    assert all((ref.get("source_id"), ref.get("locator"), ref.get("quote")) in normalized_keys
               for ref in expected_extra_refs)
    assert any("근거 역할 excluded_refs" in row["detail_text"] for row in rows)

    data, _, _ = build_output(
        {"kind": "slides", "id": "real-draft", "revision": 3,
         "template_version": "template-sha", "rows": rows},
        TEMPLATES / "Review_Deck_Template.pptx",
    )
    deck = Presentation(BytesIO(data))
    assert len(deck.slides) == 5
    for row, slide in zip(rows, deck.slides):
        visible = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
        notes = slide.notes_slide.notes_text_frame.text
        assert row["current"] in visible
        assert row["detail_text"] in notes
        for ref in row["source_refs"]:
            assert (ref.get("filename") or ref.get("source_path")) in notes
            assert str(ref.get("locator")) in notes
            assert str(ref.get("quote")) in notes


def test_slide_geometry_fonts_compact_links_and_metadata_are_readable():
    refs = [source(index) for index in range(1, 7)]
    refs[0]["filename"] = "W" * 240 + ".pdf"
    rows = build_slide_rows([])
    rows[0].update({
        "current": "Reviewed body retained exactly.",
        "decision": "REVIEW_REQUIRED",
        "rationale": "Reviewer rationale retained exactly.",
        "missing_information": ["Confirm final permit owner."],
        "detail_text": "Full generated comparison detail.",
        "excerpted": False,
        "source_refs": refs,
        "current_refs": refs,
        "historical_refs": [],
    })
    data, _, _ = build_output(
        {"kind": "slides", "id": "draft-9", "revision": 4,
         "template_version": "abc123", "rows": rows},
        TEMPLATES / "Review_Deck_Template.pptx",
    )
    deck = Presentation(BytesIO(data))
    assert deck.core_properties.author == "RE:Build Agent"
    assert deck.core_properties.last_modified_by == "RE:Build Agent"
    assert deck.core_properties.subject == "RE:Build Agent"
    assert "draft-9" in (deck.core_properties.comments or "")
    assert "abc123" in (deck.core_properties.comments or "")

    first = deck.slides[0]
    body = next(shape for shape in first.shapes if shape.name == "instructions")
    sources = next(shape for shape in first.shapes if shape.name == "source_fields")
    assert body.top >= Inches(1.4) and body.height >= Inches(3.4)
    assert sources.top >= body.top + body.height
    assert sources.top + sources.height < deck.slide_height
    assert body.text_frame.word_wrap is True
    assert all(run.font.size >= Pt(18) for p in body.text_frame.paragraphs for run in p.runs)
    assert "Reviewed body retained exactly." in body.text
    assert "Reviewer rationale retained exactly." in body.text
    assert "근거 6건" in sources.text and "발표자 노트" in sources.text
    assert all(len(paragraph.text) <= 72 for paragraph in sources.text_frame.paragraphs[1:])
    assert len([run for p in sources.text_frame.paragraphs for run in p.runs if run.hyperlink.address]) == 3
    notes = first.notes_slide.notes_text_frame.text
    assert "Full generated comparison detail." in notes
    assert refs[0]["filename"] in notes
    for ref in refs:
        assert ref["sha256"] in notes
        assert ref["quote"] in notes

    with ZipFile(BytesIO(data)) as package:
        app = ET.fromstring(package.read("docProps/app.xml"))
    ns = {"ap": "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"}
    assert app.findtext("ap:Application", namespaces=ns) == "RE:Build Agent"
    assert app.findtext("ap:Slides", namespaces=ns) == "5"
    assert app.findtext("ap:Notes", namespaces=ns) == "5"


def test_slide_export_rejects_oversize_reviewed_text_without_truncation():
    rows = build_slide_rows([])
    rows[0]["current"] = "\n".join(["검토자가 입력한 매우 긴 문장" * 4] * 12)
    with pytest.raises(ValueError, match="SLIDE_CONTENT_TOO_LONG"):
        build_output({"kind": "slides", "rows": rows}, TEMPLATES / "Review_Deck_Template.pptx")


def test_workbook_copy_has_product_identity_and_provenance(tmp_path):
    data, _, _ = build_output(
        {"kind": "itb", "id": "draft-xlsx", "revision": 2,
         "template_version": "template-xlsx-sha", "rows": []},
        TEMPLATES / "ITB_Analysis_Template.xlsx",
    )
    workbook = load_workbook(BytesIO(data))
    assert workbook.properties.creator == "RE:Build Agent"
    assert workbook.properties.lastModifiedBy == "RE:Build Agent"
    assert workbook.properties.subject == "RE:Build Agent"
    assert "draft-xlsx" in (workbook.properties.description or "")
    assert "template-xlsx-sha" in (workbook.properties.description or "")

def test_preserves_uploaded_workbook_title_and_rejects_wrong_slide_canvas(tmp_path):
    workbook = load_workbook(TEMPLATES / "ITB_Analysis_Template.xlsx")
    workbook.properties.title = "Browser uploaded immutable template version"
    custom_xlsx = tmp_path / "custom.xlsx"
    workbook.save(custom_xlsx)
    data, _, _ = build_output({"kind": "itb", "rows": []}, custom_xlsx)
    assert load_workbook(BytesIO(data)).properties.title == "Browser uploaded immutable template version"

    deck = Presentation(TEMPLATES / "Review_Deck_Template.pptx")
    deck.slide_width = Inches(7.5)
    deck.slide_height = Inches(10)
    portrait = tmp_path / "portrait.pptx"
    deck.save(portrait)
    with pytest.raises(ValueError, match="UNSUPPORTED_SLIDE_DIMENSIONS"):
        build_output({"kind": "slides", "rows": build_slide_rows([])}, portrait)


def test_long_provenance_is_lossless_in_notes_and_bounded_in_core_metadata():
    draft_id = "12345678-1234-1234-1234-123456789012"
    template_hash = "a" * 64
    data, _, _ = build_output(
        {"kind": "slides", "id": draft_id, "revision": 123,
         "template_version": template_hash, "title": "Reviewed deck",
         "rows": build_slide_rows([])},
        TEMPLATES / "Review_Deck_Template.pptx",
    )
    deck = Presentation(BytesIO(data))
    assert len(deck.core_properties.comments or "") <= 255
    notes = deck.slides[0].notes_slide.notes_text_frame.text
    assert draft_id in notes and template_hash in notes
