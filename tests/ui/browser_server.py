"""TEST_STORAGE_INJECTED isolated browser fixture, no production fallback."""
import sys, runpy
from pathlib import Path
ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT))
import backend.config
backend.config.load_environment=lambda *a,**k:None
from backend.server import create_app,Context
from backend.config import Settings
from fastapi import HTTPException
from fastapi.responses import JSONResponse
import uvicorn
fixtures=runpy.run_path(str(ROOT/'tests/test_api.py'))
from docx import Document
fixture=Document()
fixture.add_paragraph('Notice period: 17 calendar days.')
fixture.add_paragraph('Employer: Browser QA unseen synthetic contract.')
fixture.save(ROOT/'ops/runtime/browser-new.docx')
import openpyxl
version=openpyxl.load_workbook(ROOT/'data/REBUILD_INPUT_v1/REBUILD_INPUT_v1/03_OUTPUT_TEMPLATES/ITB_Analysis_Template.xlsx')
version.properties.title='Browser uploaded immutable template version'
version.save(ROOT/'ops/runtime/browser-template.xlsx')
records,files={},{}
owners={'browser-alice':'11111111-1111-4111-8111-111111111111','browser-bob':'22222222-2222-4222-8222-222222222222'}
def context(token):
    if token not in owners: raise HTTPException(401)
    return Context({'id':owners[token],'email':token+'@example.invalid'},fixtures['FakeStorage'](owners[token],records,files),fixtures['FakeAuth'](),token)
def offline_ai(*args,**kwargs):
    from backend.ai import AIError
    raise AIError('CODEX_LOGIN_REQUIRED')
app=create_app(Settings('https://wsziosnttnxefgfbgpeq.supabase.co','test-publishable','local',approval_signing_key='isolated-browser-integrity-key-not-production'),context_factory=context,ai_bridge=offline_ai)
@app.middleware('http')
async def synthetic_login(request,call_next):
    if request.url.path=='/api/auth/login':
        body=await request.json()
        if body=={'email':'browser-alice@example.invalid','password':'synthetic-test-only'}:
            return JSONResponse({'access_token':'browser-alice','test_boundary':'TEST_STORAGE_INJECTED'})
        return JSONResponse({'detail':{'code':'LOGIN_FAILED','message':'Synthetic test credentials rejected'}},401)
    response=await call_next(request)
    response.headers['X-Test-Boundary']='TEST_STORAGE_INJECTED'
    return response
if __name__=='__main__': uvicorn.run(app,host='127.0.0.1',port=8782,access_log=False)
