"""Verify downloaded browser output, TEST_STORAGE_INJECTED only."""
from pathlib import Path
import json
from datetime import datetime, timezone
from pptx import Presentation
import openpyxl
root=Path(__file__).resolve().parents[2]
workbook=openpyxl.load_workbook(root/'ops/runtime/browser-result.xlsx')
values='\n'.join(str(cell.value) for sheet in workbook for row in sheet for cell in row if cell.value is not None)
checks={
    'new_condition_retained':'17 calendar days' in values,
    'historical_condition_retained':'28 calendar days' in values,
    'reviewer_rationale_retained':'Browser reviewer verified the uploaded notice period.' in values,
    'new_source_retained':'browser-new.docx' in values,
    'historical_source_retained':'browser-history.txt' in values,
    'source_links_present':any(cell.hyperlink for sheet in workbook for row in sheet for cell in row),
}
risk=openpyxl.load_workbook(root/'ops/runtime/browser-risk.xlsx')
risk_values='\n'.join(str(cell.value) for sheet in risk for row in sheet for cell in row if cell.value is not None)
deck=Presentation(root/'ops/runtime/browser-committee.pptx')
deck_values='\n'.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
checks.update({
    'risk_reviewed_mitigation_retained':'Browser risk mitigation reviewed.' in risk_values,
    'risk_new_condition_retained':'17 calendar days' in risk_values,
    'risk_historical_condition_retained':'28 calendar days' in risk_values,
    'committee_five_actual_slides':len(deck.slides)==5,
    'committee_edited_title_retained':'Browser committee reviewed title' in deck_values,
    'committee_edited_body_retained':'Browser committee reviewed body.' in deck_values,
    'committee_reviewed_decision_retained':'Browser committee decision reviewed.' in deck_values,
    'committee_new_condition_retained':'17 calendar days' in deck_values,
    'committee_source_retained':'browser-new.docx' in deck_values,
})
custom=openpyxl.load_workbook(root/'ops/runtime/browser-custom-template-result.xlsx')
checks['uploaded_template_metadata_preserved']=custom.properties.title=='Browser uploaded immutable template version'
report={'boundary':'TEST_STORAGE_INJECTED','checks':checks,'sheets':workbook.sheetnames}
(root/'ops/runtime/browser-workbook-report.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
stamp=datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ')
(root/('ops/runtime/browser-workbook-'+stamp+'.json')).write_text(json.dumps(report,indent=2),encoding='utf-8')
print(json.dumps(report))
assert all(checks.values()),[key for key,passed in checks.items() if not passed]
