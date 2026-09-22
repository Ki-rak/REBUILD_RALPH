"""Byte-oriented document extraction with source-locatable blocks."""
from __future__ import annotations

import csv
import hashlib
import html
import io
import json
import re
import zipfile
from datetime import date, datetime, time
from email import policy
from email.parser import BytesParser
from pathlib import Path, PurePosixPath
from typing import Any, Callable

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
import pdfplumber
from pptx import Presentation


_SUPPORTED = {".docx", ".pptx", ".xlsx", ".pdf", ".md", ".txt", ".csv", ".eml"}
_OFFICE_ARCHIVES = {".docx", ".pptx", ".xlsx"}
MAX_ARCHIVE_ENTRIES = 10_000
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 100 * 1024 * 1024


class ArchiveSafetyError(ValueError):
    """Raised before parsing an unsafe Office ZIP container."""


def _validate_office_archive(data: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > MAX_ARCHIVE_ENTRIES:
            raise ArchiveSafetyError(f"Office archive entry count exceeds {MAX_ARCHIVE_ENTRIES}")
        total = 0
        for entry in entries:
            total += entry.file_size
            if total > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                raise ArchiveSafetyError(f"Office archive uncompressed size exceeds {MAX_ARCHIVE_UNCOMPRESSED_BYTES} bytes")
            path = PurePosixPath(entry.filename.replace("\\", "/"))
            if path.is_absolute() or ".." in path.parts:
                raise ArchiveSafetyError("Office archive contains an unsafe entry path")
            if entry.flag_bits & 0x1:
                raise ArchiveSafetyError("Encrypted Office archive entries are not supported")
_UNIT_RE = re.compile(
    r"(?<!\w)[+-]?(?:\d{1,3}(?:,\d{3})*|\d+)(?:\.\d+)?\s*"
    r"(?P<unit>m(?:3|³)/(?:day|d)|m(?:2|²)|m(?:3|³)|mm|cm|km|kg|kN|MPa|%|calendar\s+days?|days?|hours?)\b",
    re.IGNORECASE,
)
_REVISION_RE = re.compile(r"\b(?:revision|rev\.?)\s*[:#-]?\s*(?P<revision>[A-Za-z]?\d{1,3}|[A-Za-z])\b", re.IGNORECASE)
_STATUS_PATTERNS = (
    ("FOR REVIEW", re.compile(r"\bFOR\s+REVIEW\b", re.IGNORECASE)),
    ("APPROVED", re.compile(r"\bAPPROVED\b", re.IGNORECASE)),
    ("ISSUED FOR CONSTRUCTION", re.compile(r"\bISSUED\s+FOR\s+CONSTRUCTION\b", re.IGNORECASE)),
    ("DRAFT", re.compile(r"\bDRAFT\b", re.IGNORECASE)),
    ("SUPERSEDED", re.compile(r"\bSUPERSEDED\b", re.IGNORECASE)),
)


def _decode(data: bytes) -> tuple[str, str]:
    for encoding in ("utf-8-sig", "utf-8", "cp949"):
        try:
            return data.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace"), "utf-8-replacement"


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _unit(text: str) -> str | None:
    match = _UNIT_RE.search(text)
    return re.sub(r"\s+", " ", match.group("unit")) if match else None


def _block(
    blocks: list[dict[str, Any]], document_id: str, text: str, locator_type: str, locator: str, **extra: Any
) -> dict[str, Any]:
    item = {
        "id": f"{document_id}:block:{len(blocks) + 1}",
        "text": text,
        "locator_type": locator_type,
        "locator": locator,
        "original_value": extra.pop("original_value", text if text else None),
        "unit": extra.pop("unit", _unit(text)),
        "formula": extra.pop("formula", None),
        "cached_value": extra.pop("cached_value", None),
        "extraction_status": extra.pop("extraction_status", "EXTRACTED"),
    }
    item.update(extra)
    blocks.append(item)
    return item


def _markdown_table(rows: list[list[Any]]) -> str:
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    normalized = [["" if value is None else str(value).replace("|", "\\|").replace("\n", " ") for value in row] + [""] * (width - len(row)) for row in rows]
    header = normalized[0]
    return "\n".join(("| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * width) + " |", *("| " + " | ".join(row) + " |" for row in normalized[1:])))


def _parse_text(data: bytes, document_id: str, extension: str) -> tuple[list[dict[str, Any]], str, dict[str, Any]]:
    text, encoding = _decode(data)
    blocks: list[dict[str, Any]] = []
    for number, line in enumerate(text.splitlines(), 1):
        if line.strip():
            _block(blocks, document_id, line, "text_line", f"line {number}")
    return blocks, text, {"parser": extension.lstrip("."), "encoding": encoding}


def _parse_csv(data: bytes, document_id: str) -> tuple[list[dict[str, Any]], str, dict[str, Any]]:
    text, encoding = _decode(data)
    rows = [list(row) for row in csv.reader(io.StringIO(text))]
    blocks: list[dict[str, Any]] = []
    _block(blocks, document_id, "\n".join(", ".join(row) for row in rows), "csv_table", "rows 1-%d" % len(rows), table=rows)
    for row_number, row in enumerate(rows, 1):
        for column_number, value in enumerate(row, 1):
            if value != "":
                _block(blocks, document_id, value, "csv_cell", f"row {row_number}, column {column_number}", original_value=value)
    return blocks, _markdown_table(rows), {"parser": "csv", "encoding": encoding, "row_count": len(rows)}


def _parse_eml(data: bytes, document_id: str) -> tuple[list[dict[str, Any]], str, dict[str, Any]]:
    message = BytesParser(policy=policy.default).parsebytes(data)
    blocks: list[dict[str, Any]] = []
    markdown: list[str] = []
    for header in ("From", "To", "Cc", "Date", "Subject"):
        value = message.get(header)
        if value:
            rendered = str(value)
            _block(blocks, document_id, rendered, "eml_header", f"header {header}")
            markdown.append(f"**{header}:** {rendered}")
    body = message.get_body(preferencelist=("plain", "html")) if message.is_multipart() else message
    content = body.get_content() if body is not None else ""
    if getattr(body, "get_content_type", lambda: "")() == "text/html":
        content = html.unescape(re.sub(r"<[^>]+>", " ", content))
    for number, line in enumerate(str(content).splitlines(), 1):
        if line.strip():
            _block(blocks, document_id, line, "eml_body_line", f"body line {number}")
    if content:
        markdown.extend(("", str(content)))
    return blocks, "\n".join(markdown), {"parser": "email", "content_type": message.get_content_type()}


def _parse_docx(data: bytes, document_id: str) -> tuple[list[dict[str, Any]], str, dict[str, Any]]:
    document = Document(io.BytesIO(data))
    blocks: list[dict[str, Any]] = []
    markdown: list[str] = []
    for number, paragraph in enumerate(document.paragraphs, 1):
        text = paragraph.text.strip()
        if text:
            _block(blocks, document_id, text, "docx_paragraph", f"paragraph {number}")
            markdown.append(text)
    for table_number, table in enumerate(document.tables, 1):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        _block(blocks, document_id, "\n".join(" | ".join(row) for row in rows), "docx_table", f"table {table_number}", table=rows)
        for row_number, row in enumerate(rows, 1):
            for column_number, value in enumerate(row, 1):
                if value.strip():
                    _block(blocks, document_id, value, "docx_table_cell", f"table {table_number}, row {row_number}, column {column_number}", original_value=value)
        markdown.extend(("", _markdown_table(rows)))
    core = document.core_properties
    return blocks, "\n\n".join(markdown), {"parser": "python-docx", "title": core.title or None, "author": core.author or None, "paragraph_count": len(document.paragraphs), "table_count": len(document.tables)}


def _parse_pptx(data: bytes, document_id: str) -> tuple[list[dict[str, Any]], str, dict[str, Any]]:
    presentation = Presentation(io.BytesIO(data))
    blocks: list[dict[str, Any]] = []
    markdown: list[str] = []
    for slide_number, slide in enumerate(presentation.slides, 1):
        markdown.append(f"## Slide {slide_number}")
        for shape_number, shape in enumerate(slide.shapes, 1):
            shape_name = shape.name or f"shape {shape_number}"
            locator = f"slide {slide_number}, shape {shape_number} ({shape_name})"
            if getattr(shape, "has_table", False):
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                _block(blocks, document_id, "\n".join(" | ".join(row) for row in rows), "pptx_table", locator, table=rows)
                for row_number, row in enumerate(rows, 1):
                    for column_number, value in enumerate(row, 1):
                        if value.strip():
                            _block(blocks, document_id, value, "pptx_table_cell", f"{locator}, row {row_number}, column {column_number}", original_value=value)
                markdown.append(_markdown_table(rows))
            elif getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if text:
                    _block(blocks, document_id, text, "pptx_shape", locator)
                    markdown.append(text)
    return blocks, "\n\n".join(markdown), {"parser": "python-pptx", "slide_count": len(presentation.slides)}


def _parse_xlsx(data: bytes, document_id: str) -> tuple[list[dict[str, Any]], str, dict[str, Any]]:
    formulas = load_workbook(io.BytesIO(data), data_only=False, read_only=False)
    cached = load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    blocks: list[dict[str, Any]] = []
    markdown: list[str] = []
    try:
        for sheet in formulas.worksheets:
            cached_sheet = cached[sheet.title]
            markdown.append(f"## {sheet.title}")
            table_rows: list[list[Any]] = []
            for row in sheet.iter_rows():
                rendered_row: list[Any] = []
                for cell in row:
                    value = _json_value(cell.value)
                    rendered_row.append(value)
                    if cell.value is None:
                        continue
                    formula = cell.value if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")) else None
                    cache_value = _json_value(cached_sheet[cell.coordinate].value) if formula else None
                    extra: dict[str, Any] = {"original_value": value, "formula": formula, "cached_value": cache_value}
                    if formula and cache_value is None:
                        extra["cached_value_reason"] = "No cached formula result stored in source workbook"
                    _block(blocks, document_id, str(value), "xlsx_cell", f"{sheet.title}!{cell.coordinate}", **extra)
                if any(value is not None for value in rendered_row):
                    table_rows.append(rendered_row)
            if table_rows:
                _block(blocks, document_id, f"Worksheet {sheet.title}", "xlsx_sheet", sheet.title, table=table_rows, original_value=None)
                markdown.append(_markdown_table(table_rows))
        return blocks, "\n\n".join(markdown), {"parser": "openpyxl", "sheets": formulas.sheetnames, "calculation_mode": getattr(formulas.calculation, "calcMode", None)}
    finally:
        formulas.close()
        cached.close()


def _parse_pdf(data: bytes, document_id: str) -> tuple[list[dict[str, Any]], str, dict[str, Any]]:
    reader = PdfReader(io.BytesIO(data), strict=False)
    blocks: list[dict[str, Any]] = []
    markdown: list[str] = []
    extracted_characters = 0
    for page_number, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        status = "EXTRACTED" if text else "OCR_REQUIRED"
        _block(blocks, document_id, text, "pdf_page", f"page {page_number}", extraction_status=status, original_value=text or None)
        if text:
            extracted_characters += len(text)
            markdown.extend((f"## Page {page_number}", text))
    # Append structured evidence after every legacy page block so existing page
    # SourceRefs retain their block IDs. Coordinates use points from the top left.
    table_count = 0
    with pdfplumber.open(io.BytesIO(data)) as document:
        for page_number, page in enumerate(document.pages, 1):
            for table_number, table in enumerate(page.find_tables(), 1):
                rows = table.extract()
                if not rows:
                    continue
                table_count += 1
                locator = f"page {page_number}, table {table_number}"
                _block(blocks, document_id, "\n".join(" | ".join(value or "" for value in row) for row in rows),
                       "pdf_table", locator, table=rows, original_value=None,
                       page_number=page_number, table_number=table_number,
                       bbox=list(table.bbox), coordinate_system="pdf_points_top_left")
                for row_number, (values, geometry) in enumerate(zip(rows, table.rows), 1):
                    for column_number, value in enumerate(values, 1):
                        bbox = geometry.cells[column_number - 1]
                        _block(blocks, document_id, value or "", "pdf_table_cell",
                               f"{locator}, row {row_number}, column {column_number}",
                               original_value=value, page_number=page_number, table_number=table_number,
                               row_number=row_number, column_number=column_number,
                               cell_address=f"R{row_number}C{column_number}",
                               bbox=list(bbox) if bbox is not None else None,
                               coordinate_system="pdf_points_top_left",
                               merged_or_missing=bbox is None)
                markdown.extend((f"### Page {page_number} · Table {table_number}", _markdown_table(rows)))
    return blocks, "\n\n".join(markdown), {"parser": "pypdf", "table_parser": "pdfplumber",
        "table_detection": "ruled_geometry", "table_count": table_count,
        "page_count": len(reader.pages), "ocr_required": extracted_characters == 0}


def _detect_metadata(text: str) -> tuple[str | None, str | None]:
    revision_match = _REVISION_RE.search(text)
    revision = f"Rev{revision_match.group('revision')}" if revision_match else None
    approval = next((label for label, pattern in _STATUS_PATTERNS if pattern.search(text)), None)
    return revision, approval


def extract_document(data: bytes, filename: str, document_id: str, project_id: str) -> dict[str, Any]:
    """Extract supported document bytes without consulting paths or project-specific fixtures."""
    extension = Path(filename).suffix.lower()
    result: dict[str, Any] = {
        "id": document_id,
        "project_id": project_id,
        "filename": filename,
        "sha256": hashlib.sha256(data).hexdigest(),
        "revision": None,
        "approval_status": None,
        "storage_path": None,
        "extraction_status": "UNSUPPORTED",
        "blocks": [],
        "markdown": "",
        "metadata": {"extension": extension, "byte_size": len(data)},
    }
    if extension not in _SUPPORTED:
        result["metadata"].update({"error_type": "UnsupportedFormat", "supported_extensions": sorted(_SUPPORTED)})
        return result
    parsers: dict[str, Callable[..., tuple[list[dict[str, Any]], str, dict[str, Any]]]] = {
        ".docx": _parse_docx, ".pptx": _parse_pptx, ".xlsx": _parse_xlsx, ".pdf": _parse_pdf,
        ".md": lambda value, doc: _parse_text(value, doc, ".md"), ".txt": lambda value, doc: _parse_text(value, doc, ".txt"),
        ".csv": _parse_csv, ".eml": _parse_eml,
    }
    try:
        if extension in _OFFICE_ARCHIVES:
            _validate_office_archive(data)
        blocks, markdown, metadata = parsers[extension](data, document_id)
        revision, approval = _detect_metadata("\n".join(block.get("text", "") for block in blocks))
        status = "OCR_REQUIRED" if metadata.get("ocr_required") else ("EXTRACTED" if blocks else "EMPTY")
        result.update({"revision": revision, "approval_status": approval, "extraction_status": status, "blocks": blocks, "markdown": markdown})
        result["metadata"].update(metadata)
        result["metadata"].update({"revision_source": "content" if revision else None, "approval_status_source": "content" if approval else None})
    except Exception as exc:
        result["extraction_status"] = "ERROR"
        result["metadata"].update({"error_type": type(exc).__name__, "error": str(exc)[:500]})
    return result
