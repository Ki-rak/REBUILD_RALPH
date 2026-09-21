"""Owner-scoped immutable copies of reusable Office template versions."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import hmac
import json
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation

from .extraction import extract_document
from .storage import ConflictError

FIELDS = ("template_id", "filename", "sha256", "storage_path", "origin", "created_at")


class TemplateError(ValueError):
    pass


class TemplateCatalog:
    def __init__(self, store, owner, signing_key, provided):
        self.store, self.owner, self.key, self.provided = store, owner, signing_key, provided

    def _seal(self, item):
        message = json.dumps({"owner": self.owner, "template": {k: item.get(k) for k in FIELDS}},
                             sort_keys=True, ensure_ascii=False, separators=(",", ":")).encode()
        return hmac.new(self.key.encode(), message, "sha256").hexdigest()

    def _verify(self, record):
        item = deepcopy(record["payload"])
        signature = item.get("signature")
        if not isinstance(signature, str) or not hmac.compare_digest(signature, self._seal(item)):
            raise TemplateError("TEMPLATE_INTEGRITY_FAILED")
        return item

    def _base(self, tid):
        if tid not in self.provided:
            raise TemplateError("TEMPLATE_NOT_FOUND")
        info, path = self.provided[tid]
        return deepcopy(info), path.read_bytes()

    def versions(self, tid):
        base, _ = self._base(tid)
        versions = {base["sha256"]: {**base, "version": base["sha256"], "origin": "provided", "created_at": None}}
        for record in self.store.list("template"):
            item = self._verify(record)
            if item["template_id"] == tid:
                versions[item["sha256"]] = {**base, **item, "version": item["sha256"]}
        return sorted(versions.values(), key=lambda item: item.get("created_at") or "", reverse=True)

    def resolve(self, tid, version=None, persist=False):
        base, original = self._base(tid)
        selected = version or base["sha256"]
        record = self.store.get(f"template:{tid}:{selected}")
        if record:
            item = self._verify(record)
            if item["template_id"] != tid or item["sha256"] != selected:
                raise TemplateError("TEMPLATE_INTEGRITY_FAILED")
            data = self.store.download_bytes(item["storage_path"])
            if sha256(data).hexdigest() != selected:
                raise TemplateError("TEMPLATE_INTEGRITY_FAILED")
            return {**base, **item, "version": selected}, data
        if selected != base["sha256"]:
            raise TemplateError("TEMPLATE_VERSION_NOT_FOUND")
        if persist:
            self._save(tid, base["filename"], original, "provided")
            return self.resolve(tid, selected)
        return {**base, "version": selected, "origin": "provided"}, original

    def _save(self, tid, filename, data, origin):
        from datetime import datetime, timezone
        digest = sha256(data).hexdigest()
        eid = f"template:{tid}:{digest}"
        if self.store.get(eid):
            return self.resolve(tid, digest)[0]
        path = f"{self.owner}/templates/{tid}/{digest}/{filename}"
        try:
            self.store.upload_bytes(path, data, "application/octet-stream")
        except ConflictError:
            # A prior interrupted registration may have stored bytes before its DB row.
            if sha256(self.store.download_bytes(path)).hexdigest() != digest:
                raise TemplateError("TEMPLATE_INTEGRITY_FAILED") from None
        item = {"template_id": tid, "filename": filename, "sha256": digest, "size": len(data),
                "storage_path": path, "origin": origin,
                "created_at": datetime.now(timezone.utc).isoformat()}
        item["signature"] = self._seal(item)
        self.store.save("template", eid, None, item)
        return {**self.provided[tid][0], **item, "version": digest}

    def register(self, tid, filename, data):
        base, original = self._base(tid)
        filename = filename.replace("\\", "/").rsplit("/", 1)[-1]
        if not filename or len(filename) > 240 or any(ord(c) < 32 for c in filename):
            raise TemplateError("INVALID_TEMPLATE_FILENAME")
        if Path(filename).suffix.lower() != Path(base["filename"]).suffix.lower():
            raise TemplateError("TEMPLATE_TYPE_MISMATCH")
        result = extract_document(data, filename, "template-validation", "template-validation")
        if result.get("extraction_status") != "EXTRACTED":
            raise TemplateError("INVALID_TEMPLATE")
        try:
            if tid in {"itb", "risk"}:
                sheet_name, width = ("ITB", 8) if tid == "itb" else ("RiskOutput", 10)
                with BytesIO(original) as source, BytesIO(data) as candidate:
                    before = load_workbook(source, read_only=True)
                    after = load_workbook(candidate, read_only=True)
                    try:
                        headers = [before[sheet_name].cell(5, c).value for c in range(1, width+1)]
                        if [after[sheet_name].cell(5, c).value for c in range(1, width+1)] != headers:
                            raise TemplateError("TEMPLATE_COLUMNS_CHANGED")
                    finally:
                        before.close()
                        after.close()
            elif tid == "slides":
                deck = Presentation(BytesIO(data))
                provided_deck = Presentation(BytesIO(original))
                if (deck.slide_width, deck.slide_height) != (provided_deck.slide_width, provided_deck.slide_height):
                    raise TemplateError("SLIDE_TEMPLATE_SIZE_UNSUPPORTED")
                if len(deck.slides) != 5:
                    raise TemplateError("FIVE_SLIDE_TEMPLATE_REQUIRED")
                for slide in deck.slides:
                    texts = [shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False)]
                    if not any("[신규 프로젝트 내용을 입력]" in text for text in texts) or not any("[원문 문서번호" in text for text in texts):
                        raise TemplateError("TEMPLATE_PLACEHOLDERS_REQUIRED")
        except (KeyError, ValueError) as error:
            if isinstance(error, TemplateError):
                raise
            raise TemplateError("INVALID_TEMPLATE") from None
        self._save(tid, base["filename"], original, "provided")
        return self._save(tid, filename, data, "uploaded")