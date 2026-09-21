"""Editable Office outputs built from approved draft data and provided templates."""
from __future__ import annotations

import io
import json
from copy import copy
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import parse_qs, urlparse
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.workbook.properties import CalcProperties
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt
from .drafts import SLIDE_IDS, SLIDE_VISIBLE_SOURCE_LIMIT, slide_visible_text, validate_slide_rows

_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
_REVIEW = "REVIEW_REQUIRED"
_PRODUCT = "RE:Build Agent"
_SLIDE_WIDTH = 12192000
_SLIDE_HEIGHT = 6858000


def _first(mapping: dict[str, Any], *keys: str, default: Any = None) -> Any:
    for key in keys:
        value = mapping.get(key)
        if value is not None:
            return value
    return default


def _items(draft: dict[str, Any], *keys: str) -> list[Any]:
    for key in keys:
        value = draft.get(key)
        if isinstance(value, list):
            return value
    return []


def _text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, (list, tuple, set)):
        rendered = [str(entry) for entry in value if entry not in (None, "")]
        return "\n".join(rendered) or None
    if isinstance(value, dict):
        rendered = [f"{key}: {entry}" for key, entry in value.items() if entry not in (None, "")]
        return "\n".join(rendered) or None
    return str(value)


def _refs(item: dict[str, Any]) -> list[dict[str, Any]]:
    value = _first(item, "source_refs", "sources", "evidence", default=[])
    if isinstance(value, dict):
        return [value]
    if isinstance(value, list):
        return [entry for entry in value if isinstance(entry, dict)]
    source = item.get("source_ref")
    return [source] if isinstance(source, dict) else []


def _render_refs(refs: Iterable[dict[str, Any]]) -> str:
    rendered: list[str] = []
    for ref in refs:
        path = _first(ref, "source_path", "filename", "path", "document_id", default="source")
        locator = _first(ref, "locator", "location", default="locator unavailable")
        quote = str(ref.get("quote") or "").strip()
        line = f"[{path} | {locator}]"
        if quote:
            line += f" {quote}"
        rendered.append(line)
    return "\n".join(rendered) if rendered else _REVIEW


def _safe_url(value: Any) -> str | None:
    if not isinstance(value, str):
        return None
    parsed = urlparse(value)
    if parsed.scheme not in {"http", "https"} or not parsed.netloc or parsed.username or parsed.password:
        return None
    forbidden = {"token", "access_token", "auth", "authorization", "api_key", "apikey"}
    if forbidden.intersection(key.casefold() for key in parse_qs(parsed.query)):
        return None
    return value


def _ref_key(ref: dict[str, Any]) -> tuple[Any, ...]:
    return (ref.get("source_id"), ref.get("document_id"), ref.get("sha256"), ref.get("locator"), ref.get("quote"))


def _collect_refs(value: Any) -> list[dict[str, Any]]:
    found: list[dict[str, Any]] = []
    if isinstance(value, dict):
        if (value.get("source_id") or value.get("document_id")) and value.get("locator") is not None:
            found.append(value)
        else:
            for nested in value.values():
                found.extend(_collect_refs(nested))
    elif isinstance(value, list):
        for nested in value:
            found.extend(_collect_refs(nested))
    unique = {}
    for ref in found:
        unique.setdefault(_ref_key(ref), ref)
    return list(unique.values())


def _write_value(cell: Any, value: Any) -> None:
    """Preserve source/user strings literally; only explicit generated formulas execute."""
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def _add_sources_sheet(workbook: Any, draft: dict[str, Any]) -> dict[tuple[Any, ...], int]:
    if "Sources" in workbook.sheetnames:
        del workbook["Sources"]
    sheet = workbook.create_sheet("Sources")
    headers = ["Source ID", "Document ID", "Source path", "SHA-256", "Revision", "Approval status", "Locator type", "Locator", "Quote", "Open source"]
    for column, header in enumerate(headers, 1):
        cell = sheet.cell(1, column, header)
        cell.font = Font(bold=True)
    rows = {}
    for row, ref in enumerate(_collect_refs(draft), 2):
        values = [ref.get("source_id"), ref.get("document_id"), _first(ref, "source_path", "filename"), ref.get("sha256"),
                  ref.get("revision"), ref.get("approval_status"), ref.get("locator_type"), ref.get("locator"), ref.get("quote"), "원문 열기"]
        for column, value in enumerate(values, 1):
            _write_value(sheet.cell(row, column), value)
            sheet.cell(row, column).alignment = Alignment(vertical="top", wrap_text=True)
        url = _safe_url(ref.get("url"))
        if url:
            sheet.cell(row, 10).hyperlink = url
            sheet.cell(row, 10).font = Font(color="0563C1", underline="single")
        else:
            sheet.cell(row, 10).value = "앱 원문 링크 없음"
        rows[_ref_key(ref)] = row
    for column, width in enumerate((28, 24, 45, 68, 14, 18, 18, 26, 80, 22), 1):
        sheet.column_dimensions[sheet.cell(1, column).column_letter].width = width
    return rows


def _link_to_source(cell: Any, refs: list[dict[str, Any]], source_rows: dict[tuple[Any, ...], int]) -> None:
    if refs:
        row = source_rows.get(_ref_key(refs[0]))
        if row:
            cell.hyperlink = f"#'Sources'!A{row}"
            cell.font = copy(cell.font)
            cell.font = Font(name=cell.font.name, size=cell.font.sz, bold=cell.font.bold, italic=cell.font.italic,
                             color="0563C1", underline="single")


def _review(item: dict[str, Any], has_sources: bool, required_fields: Iterable[str] = ()) -> str:
    explicit = _first(item, "review_required", "review_notes", "review_status", "confirmation_required")
    missing = item.get("missing_information")
    notes: list[str] = []
    if isinstance(explicit, str) and explicit.strip():
        notes.append(explicit.strip())
    elif explicit is True:
        notes.append(_REVIEW)
    if isinstance(missing, list):
        notes.extend(str(value) for value in missing if value)
    elif isinstance(missing, str) and missing.strip():
        notes.append(missing.strip())
    absent = [field for field in required_fields if item.get(field) in (None, "", [])]
    if not has_sources:
        notes.append("source evidence missing")
    if absent:
        notes.append("missing: " + ", ".join(absent))
    if notes and not any(_REVIEW in note for note in notes):
        notes.insert(0, _REVIEW)
    return "; ".join(dict.fromkeys(notes))


def _copy_row_style(sheet: Any, source_row: int, target_row: int, width: int) -> None:
    if target_row == source_row:
        return
    sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height
    for column in range(1, width + 1):
        source = sheet.cell(source_row, column)
        target = sheet.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        target.number_format = source.number_format
        target.protection = copy(source.protection)
        target.alignment = copy(source.alignment)


def _provenance(draft: dict[str, Any], template_path: Path, existing: str | None = None) -> str:
    parts = [part for part in (existing, f"Generated by {_PRODUCT}", f"Template: {template_path.name}") if part]
    for label, keys in (
        ("Draft title", ("title",)),
        ("Draft ID", ("id", "draft_id")),
        ("Revision", ("revision",)),
        ("Template SHA-256", ("template_version", "template_sha256", "template_hash")),
    ):
        value = _first(draft, *keys)
        if value not in (None, ""):
            parts.append(f"{label}: {value}")
    return " | ".join(parts)


def _bounded_core_text(value: str, limit: int = 255) -> str:
    """Keep constrained Office core fields valid and point to the lossless copy."""
    if len(value) <= limit:
        return value
    suffix = " | Full provenance: slide notes"
    return value[:limit - len(suffix) - 1].rstrip() + "…" + suffix


def _set_workbook_metadata(workbook: Any, draft: dict[str, Any], template_path: Path) -> None:
    properties = workbook.properties
    original = " | ".join(filter(None, (
        f"Template title: {properties.title}" if properties.title else None,
        f"Template creator: {properties.creator}" if properties.creator else None,
        properties.description,
    )))
    properties.creator = _PRODUCT
    properties.lastModifiedBy = _PRODUCT
    if not properties.title:
        properties.title = f"{_PRODUCT} · {draft.get('title') or template_path.stem}"
    properties.subject = _PRODUCT
    properties.keywords = _PRODUCT
    properties.description = _provenance(draft, template_path, original)


def _save_workbook(workbook: Any) -> bytes:
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _build_itb(draft: dict[str, Any], template_path: Path) -> bytes:
    workbook = load_workbook(template_path)
    _set_workbook_metadata(workbook, draft, template_path)
    sheet = workbook["ITB"] if "ITB" in workbook.sheetnames else workbook.active
    source_rows = _add_sources_sheet(workbook, draft)
    rows = _items(draft, "items", "rows", "itb_items")
    for index, item_value in enumerate(rows, 6):
        item = item_value if isinstance(item_value, dict) else {"current_condition": str(item_value)}
        _copy_row_style(sheet, 13, index, 8)
        refs = _refs(item)
        current = _first(item, "current_condition", "new_condition", "condition", "current")
        decision = _first(item, "decision", "applicability", "judgement")
        rationale = _text(item.get("rationale"))
        judgement = "\n".join(part for part in (
            f"적용 판단: {_text(decision)}" if decision not in (None, "") else None,
            f"근거 설명: {rationale}" if rationale else None,
        ) if part)
        review_item = {**item, "current_condition": current, "decision": decision}
        values = [
            _first(item, "item_id", "id", default=f"ITEM-{index - 5:02d}"),
            _first(item, "itb_clause", "clause", "title"),
            _text(current),
            _text(_first(item, "historical_case", "past_case", "history", "past")),
            _text(_first(item, "differences", "difference")),
            judgement or None,
            _render_refs(refs),
            _review(review_item, bool(refs), ("current_condition", "decision")),
        ]
        for column, value in enumerate(values, 1):
            _write_value(sheet.cell(index, column), value)
            sheet.cell(index, column).alignment = copy(sheet.cell(index, column).alignment)
            sheet.cell(index, column).alignment = Alignment(horizontal=sheet.cell(index, column).alignment.horizontal, vertical="top", wrap_text=True)
        _link_to_source(sheet.cell(index, 7), refs, source_rows)
    for row in range(6 + len(rows), 14):
        for column in range(1, 9):
            sheet.cell(row, column).value = None

    if "Key Facts" in workbook.sheetnames:
        del workbook["Key Facts"]
    facts_sheet = workbook.create_sheet("Key Facts")
    headers = ["Fact", "Value", "Unit", "Source evidence", "Review status"]
    for column, header in enumerate(headers, 1):
        cell = facts_sheet.cell(1, column, header)
        cell.font = Font(bold=True)
    facts = draft.get("key_facts") or []
    if not facts and rows:
        facts = [{"fact": _first(item, "title", "itb_clause", "id"),
                  "value": _first(item, "current", "current_condition", "new_condition"),
                  "source_refs": item.get("current_refs") or item.get("source_refs") or [],
                  "missing_information": item.get("missing_information") or []}
                 for item in rows if isinstance(item, dict)]
    if isinstance(facts, dict):
        facts = [{"fact": key, "value": value} for key, value in facts.items()]
    for row, fact_value in enumerate(facts, 2):
        fact = fact_value if isinstance(fact_value, dict) else {"fact": str(fact_value)}
        refs = _refs(fact)
        values = [_first(fact, "fact", "name", "label"), fact.get("value"), fact.get("unit"), _render_refs(refs), _review(fact, bool(refs), ("value",))]
        for column, value in enumerate(values, 1):
            _write_value(facts_sheet.cell(row, column), value)
            facts_sheet.cell(row, column).alignment = Alignment(vertical="top", wrap_text=True)
        _link_to_source(facts_sheet.cell(row, 4), refs, source_rows)
    if not facts:
        facts_sheet.append([_REVIEW, None, None, _REVIEW, "No key facts supplied"])
    for column, width in enumerate((28, 24, 16, 70, 42), 1):
        facts_sheet.column_dimensions[facts_sheet.cell(1, column).column_letter].width = width
    return _save_workbook(workbook)


def _build_risk(draft: dict[str, Any], template_path: Path) -> bytes:
    workbook = load_workbook(template_path)
    _set_workbook_metadata(workbook, draft, template_path)
    sheet = workbook["RiskOutput"] if "RiskOutput" in workbook.sheetnames else workbook.active
    source_rows = _add_sources_sheet(workbook, draft)
    rows = _items(draft, "items", "rows", "risks", "risk_items")
    for index, item_value in enumerate(rows, 6):
        item = item_value if isinstance(item_value, dict) else {"risk": str(item_value)}
        _copy_row_style(sheet, 13, index, 10)
        refs = _refs(item)
        probability = _first(item, "probability", "likelihood")
        intensity = _first(item, "intensity", "severity", "impact_rating")
        score = item.get("score")
        narrative = " / ".join(str(value) for value in (_first(item, "cause", "rationale"), _first(item, "impact", "consequence", "current"), _text(item.get("past"))) if value not in (None, "")) or None
        review = _review(item, bool(refs), ("risk", "mitigation"))
        if probability in (None, "") or intensity in (None, ""):
            score = None
            if not review:
                review = f"{_REVIEW}; probability/intensity not established"
            elif "probability/intensity" not in review:
                review += "; probability/intensity not established"
        values = [
            _first(item, "risk_id", "item_id", "id", default=f"R-{index - 5:02d}"),
            _first(item, "risk", "title", "description"), narrative, probability, intensity, score,
            _first(item, "mitigation", "response", "action"), _first(item, "owner", "responsible"), _render_refs(refs), review or "REVIEWED",
        ]
        for column, value in enumerate(values, 1):
            _write_value(sheet.cell(index, column), value)
            sheet.cell(index, column).alignment = Alignment(vertical="top", wrap_text=True)
        _link_to_source(sheet.cell(index, 9), refs, source_rows)
        if score is None and probability not in (None, "") and intensity not in (None, ""):
            sheet.cell(index, 6).value = f'=IF(D{index}="","",IF(E{index}="","",D{index}*E{index}))'
    for row in range(6 + len(rows), 14):
        for column in range(1, 11):
            sheet.cell(row, column).value = None
    if workbook.calculation is None:
        workbook.calculation = CalcProperties()
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    return _save_workbook(workbook)


def _slide_lines(values: list[Any]) -> tuple[str, str]:
    body_lines: list[str] = []
    source_lines: list[str] = []
    if not values:
        return f"• {_REVIEW} — 근거가 있는 항목 없음", _REVIEW
    for value in values:
        item = value if isinstance(value, dict) else {"text": str(value)}
        text = _text(_first(item, "text", "title", "summary", "condition", "risk", default=_REVIEW))
        refs = _refs(item)
        review = _review(item, bool(refs))
        body_lines.append(f"• {text}" + (f" [{review}]" if review else ""))
        source_lines.append(_render_refs(refs))
    return "\n".join(body_lines), "\n".join(source_lines)


def _slide_refs(values: list[Any]) -> list[dict[str, Any]]:
    unique: dict[tuple[Any, ...], dict[str, Any]] = {}
    for value in values:
        if not isinstance(value, dict):
            continue
        for field in ("source_refs", "current_refs", "historical_refs"):
            refs = value.get(field) or []
            if isinstance(refs, dict):
                refs = [refs]
            for ref in refs:
                if isinstance(ref, dict):
                    unique.setdefault(_ref_key(ref), ref)
    return list(unique.values())


def _compact_ref(ref: dict[str, Any]) -> str:
    path = str(_first(ref, "filename", "source_path", "path", "document_id", default="source"))
    name = path.replace("\\", "/").rsplit("/", 1)[-1]
    locator = str(_first(ref, "locator", "location", default="locator unavailable"))
    label = f"{name} · {locator}"
    return label if len(label) <= 72 else label[:69].rstrip() + " …"


def _style_frame(shape: Any, font_size: int, *, top_margin: float = 0.05) -> None:
    frame = shape.text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.08)
    frame.margin_top = Inches(top_margin)
    frame.margin_bottom = Inches(0.03)
    for paragraph in frame.paragraphs:
        paragraph.line_spacing = Pt(font_size)
        paragraph.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.size = Pt(font_size)


def _set_ppt_sources(shape: Any, refs: list[dict[str, Any]]) -> None:
    frame = shape.text_frame
    frame.clear()
    if not refs:
        frame.paragraphs[0].text = f"{_REVIEW} · 연결된 원문 근거 없음"
        _style_frame(shape, 9, top_margin=0.02)
        return
    frame.paragraphs[0].text = f"근거 {len(refs)}건 · 전체 SourceRef는 발표자 노트에서 확인"
    for ref in refs[:SLIDE_VISIBLE_SOURCE_LIMIT]:
        paragraph = frame.add_paragraph()
        run = paragraph.add_run()
        run.text = _compact_ref(ref)
        url = _safe_url(ref.get("url"))
        if url:
            run.hyperlink.address = url
    _style_frame(shape, 9, top_margin=0.02)


def _set_ppt_notes(slide: Any, row: dict[str, Any], refs: list[dict[str, Any]], provenance: str) -> None:
    detail = _text(row.get("detail_text")) or slide_visible_text(row) or _REVIEW
    lines = [
        _PRODUCT,
        f"슬라이드: {_text(row.get('title')) or row.get('id') or _REVIEW}",
        "",
        "검토 본문 원문",
        f"current: {_text(row.get('current'))}",
        f"decision: {_text(row.get('decision'))}",
        f"rationale: {_text(row.get('rationale'))}",
        f"missing_information: {_text(row.get('missing_information'))}",
        "",
        "출력 provenance",
        provenance,
        "",
        "전체 비교 상세",
        detail,
        "",
        f"전체 SourceRefs ({len(refs)}건)",
    ]
    if refs:
        for index, ref in enumerate(refs, 1):
            lines.append(f"SourceRef {index}")
            for key in sorted(ref):
                value = ref[key]
                rendered = json.dumps(value, ensure_ascii=False, sort_keys=True) if isinstance(value, (dict, list)) else str(value)
                lines.append(f"{key}: {rendered}")
    else:
        lines.append(_REVIEW)
    slide.notes_slide.notes_text_frame.text = "\n".join(lines)


def _set_presentation_metadata(deck: Any, draft: dict[str, Any], template_path: Path) -> str:
    properties = deck.core_properties
    original = " | ".join(filter(None, (
        f"Template title: {properties.title}" if properties.title else None,
        f"Template creator: {properties.author}" if properties.author else None,
        properties.comments,
    )))
    provenance = _provenance(draft, template_path, original)
    properties.author = _PRODUCT
    properties.last_modified_by = _PRODUCT
    properties.title = _bounded_core_text(f"{_PRODUCT} · {draft.get('title') or '심의장표'}")
    properties.subject = _PRODUCT
    properties.keywords = _PRODUCT
    properties.comments = _bounded_core_text(provenance)
    return provenance


def _patch_pptx_app_metadata(data: bytes, slide_count: int, notes_count: int) -> bytes:
    source = io.BytesIO(data)
    output = io.BytesIO()
    namespace = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
    with ZipFile(source, "r") as archive, ZipFile(output, "w") as rewritten:
        for info in archive.infolist():
            payload = archive.read(info.filename)
            if info.filename == "docProps/app.xml":
                root = ET.fromstring(payload)
                for name, value in (("Application", _PRODUCT), ("Slides", str(slide_count)), ("Notes", str(notes_count))):
                    node = root.find(f"{{{namespace}}}{name}")
                    if node is None:
                        node = ET.SubElement(root, f"{{{namespace}}}{name}")
                    node.text = value
                payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            rewritten.writestr(info, payload)
    return output.getvalue()


def _build_slides(draft: dict[str, Any], template_path: Path) -> bytes:
    deck = Presentation(template_path)
    if (deck.slide_width, deck.slide_height) != (_SLIDE_WIDTH, _SLIDE_HEIGHT):
        raise ValueError("UNSUPPORTED_SLIDE_DIMENSIONS")
    rows = _items(draft, "rows")
    canonical = [row.get("id") for row in rows] == list(SLIDE_IDS)
    if canonical:
        validate_slide_rows(rows)
    provenance = _set_presentation_metadata(deck, draft, template_path)
    sections = (
        _items(draft, "business", "overview", "business_items"),
        _items(draft, "itb", "itb_items", "conditions"),
        _items(draft, "applicability", "applicability_items", "comparisons"),
        _items(draft, "risks", "risk_items"),
        _items(draft, "unknowns", "missing_information", "decisions"),
    )
    if canonical:
        # The reviewed slide body is authoritative; do not regenerate it from topics.
        sections = tuple([{
            "text": "\n".join(filter(None, (
                _text(row.get("current")), _text(row.get("decision")),
                _text(row.get("rationale")), _text(row.get("missing_information")),
            ))),
            "source_refs": _refs(row),
        }] for row in rows)
    elif rows and not any(sections):
        # Compatibility for the original topic-based draft representation.
        business = [{"text": f"{_first(row, 'title', 'id')}: {_text(row.get('current')) or _REVIEW}", "source_refs": row.get("current_refs") or row.get("source_refs") or []} for row in rows]
        itb = [{"text": _text(row.get("current")) or f"{_first(row, 'title', 'id')}: {_REVIEW}", "source_refs": row.get("current_refs") or row.get("source_refs") or []} for row in rows]
        applicability = [{"text": "\n".join(filter(None, (_text(row.get("title")), _text(row.get("differences")), _text(row.get("decision")), _text(row.get("rationale"))))), "source_refs": row.get("source_refs") or []} for row in rows]
        risks = [{"text": "\n".join(filter(None, (_text(row.get("title")), _text(row.get("mitigation"))))), "source_refs": row.get("source_refs") or []} for row in rows]
        unknowns = [{"text": _text(entry), "source_refs": row.get("source_refs") or []} for row in rows for entry in (row.get("missing_information") or [_REVIEW])]
        sections = (business, itb, applicability, risks, unknowns)
    if len(deck.slides) < 5:
        raise ValueError("Review deck template must contain at least five slides")
    for slide_number, (slide, items) in enumerate(zip(list(deck.slides)[:5], sections), 1):
        row = rows[slide_number - 1] if canonical else None
        body, sources = (slide_visible_text(row), "") if canonical else _slide_lines(items)
        refs = _slide_refs([row]) if canonical else _slide_refs(items)
        body_set = False
        source_set = False
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            current = shape.text.strip()
            if "[신규 프로젝트 내용을 입력]" in current or current.startswith("[Enter project"):
                shape.text = body
                shape.left, shape.top = Inches(0.73), Inches(1.48)
                shape.width, shape.height = Inches(11.46), Inches(3.55)
                _style_frame(shape, 18)
                body_set = True
            elif "[원문 문서번호" in current or current.startswith("[Source"):
                shape.left, shape.top = Inches(0.73), Inches(5.15)
                shape.width, shape.height = Inches(11.46), Inches(0.95)
                _set_ppt_sources(shape, refs)
                source_set = True
        text_shapes = [shape for shape in slide.shapes if getattr(shape, "has_text_frame", False)]
        if not body_set and len(text_shapes) > 1:
            text_shapes[1].text = body
            text_shapes[1].left, text_shapes[1].top = Inches(0.73), Inches(1.48)
            text_shapes[1].width, text_shapes[1].height = Inches(11.46), Inches(3.55)
            _style_frame(text_shapes[1], 18)
        if not source_set and len(text_shapes) > 2:
            text_shapes[2].left, text_shapes[2].top = Inches(0.73), Inches(5.15)
            text_shapes[2].width, text_shapes[2].height = Inches(11.46), Inches(0.95)
            _set_ppt_sources(text_shapes[2], refs)
        title = rows[slide_number - 1].get("title") if canonical else draft.get("title") if slide_number == 1 else None
        if title and text_shapes:
            title_shape = next((shape for shape in text_shapes
                                if getattr(shape, "is_placeholder", False)
                                and getattr(shape.placeholder_format, "type", None) == 1), text_shapes[0])
            title_shape.text = str(title)
            title_shape.left, title_shape.top = Inches(0.67), Inches(0.36)
            title_shape.width, title_shape.height = Inches(11.98), Inches(0.98)
            _style_frame(title_shape, 28, top_margin=0.02)
        evidence = next((shape for shape in text_shapes if shape.name == "evidence_note"), None)
        if evidence is not None:
            evidence.left, evidence.top = Inches(0.73), Inches(6.23)
            evidence.width, evidence.height = Inches(11.46), Inches(0.34)
            evidence.text = "전체 비교 상세와 SourceRef는 발표자 노트에 보존됩니다."
            _style_frame(evidence, 10, top_margin=0.01)
        footer = next((shape for shape in text_shapes if shape.name == "footer"), None)
        if footer is not None:
            footer.text = f"{_PRODUCT} | 근거 검토용 | {slide_number}/5"
            _style_frame(footer, 10, top_margin=0.0)
        note_row = row or {"id": f"slide-{slide_number}", "title": title, "detail_text": body}
        _set_ppt_notes(slide, note_row, refs, provenance)
    output = io.BytesIO()
    deck.save(output)
    return _patch_pptx_app_metadata(output.getvalue(), 5, 5)


def build_output(draft: dict[str, Any], template_path: Path) -> tuple[bytes, str, str]:
    """Populate a provided template; approval enforcement remains the caller's responsibility."""
    template = Path(template_path)
    if not template.is_file():
        raise FileNotFoundError(template)
    kind = str(draft.get("kind") or draft.get("type") or "").strip().lower()
    name = template.name.lower()
    if not kind:
        kind = "itb" if "itb" in name else "risk" if "risk" in name else "slides" if template.suffix.lower() == ".pptx" else ""
    if kind in {"itb", "itb_analysis"}:
        if template.suffix.lower() != ".xlsx":
            raise ValueError("ITB output requires an XLSX template")
        return _build_itb(draft, template), _XLSX_MIME, ".xlsx"
    if kind in {"risk", "risk_register"}:
        if template.suffix.lower() != ".xlsx":
            raise ValueError("Risk output requires an XLSX template")
        return _build_risk(draft, template), _XLSX_MIME, ".xlsx"
    if kind in {"slides", "review_deck", "pptx"}:
        if template.suffix.lower() != ".pptx":
            raise ValueError("Slide output requires a PPTX template")
        return _build_slides(draft, template), _PPTX_MIME, ".pptx"
    raise ValueError(f"Unsupported output kind: {kind or 'unspecified'}")
