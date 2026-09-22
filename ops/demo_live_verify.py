"""Real INPUT -> real Supabase -> real product HTTP -> Office -> process restart.

Uses disposable, explicitly marked test users. No demo account is preloaded with
new-project files, and no human approval is fabricated. Rules/AI proofs are split.
"""
from __future__ import annotations
import argparse
from copy import deepcopy
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
import secrets
import socket
import subprocess
import sys
import time
from uuid import UUID, uuid4

ROOT=Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:sys.path.insert(0,str(ROOT))
import httpx
from openpyxl import load_workbook
from pptx import Presentation
from backend.config import Settings, load_environment
from backend.storage import SupabaseStore
from ops.demo import ProductAPI, DemoError, catalog, select_projects, seed_past, upload_new, read_original, write_report
from ops.runtime.supabase_live_verify import verify_live, _create_test_user, _service_headers


def require(condition, code):
    if not condition:raise DemoError(code)


def check(report, name, **details):
    report['checks'].append({'name':name,'passed':True,**details})


def office_signature(data: bytes, kind: str) -> dict:
    try:
        if kind=='slides':
            deck=Presentation(BytesIO(data))
            require(len(deck.slides)==5,'SLIDE_COUNT_MISMATCH')
            content=[{'text':[shape.text for shape in slide.shapes if hasattr(shape,'text')],
                      'notes':slide.notes_slide.notes_text_frame.text} for slide in deck.slides]
            structure={'slides':len(deck.slides)}
        else:
            book=load_workbook(BytesIO(data),data_only=False)
            content={sheet.title:[[cell.coordinate,str(cell.value),cell.data_type] for row in sheet for cell in row if cell.value is not None] for sheet in book}
            require('Sources' in content,'OFFICE_SOURCES_MISSING')
            structure={'sheets':book.sheetnames}
        encoded=json.dumps(content,ensure_ascii=False,sort_keys=True).encode('utf-8')
        return {**structure,'semantic_sha256':sha256(encoded).hexdigest()}
    except DemoError:raise
    except Exception:raise DemoError('OFFICE_REOPEN_FAILED') from None


def run_test_flow(api, past, new, output_dir: Path, report: dict) -> dict:
    output_dir.mkdir(parents=True,exist_ok=True)
    report['reviewer']='AUTOMATED_DISPOSABLE_TEST_ACCOUNT_ONLY'
    report['historical']=seed_past(api,past,approved=True)
    check(report,'approved_past_originals_imported',projects=len(past))
    snapshot={'projects':[],'documents':[],'historical_documents':[],'drafts':[]}
    for selected, imported in zip(past, report['historical']):
        originals={item['sha256']:item for item in selected['files']}
        for doc in imported['documents']:
            data=api.request('GET',f"/api/documents/{doc['id']}/original").content
            require(data==read_original(originals[doc['sha256']]),'HISTORICAL_ORIGINAL_MISMATCH')
            knowledge=api.json('GET',f"/api/documents/{doc['id']}/knowledge")
            require(knowledge.get('markdown') and knowledge.get('blocks') and knowledge.get('sidecars'),'HISTORICAL_KNOWLEDGE_MISSING')
            snapshot['historical_documents'].append({'id':doc['id'],'sha256':doc['sha256'],'sidecars':knowledge['sidecars']})
    for selected in new:
        incoming=upload_new(api,selected)
        report.setdefault('uploads',[]).append({key:value for key,value in incoming.items() if key!='documents'})
        pid=incoming['project_id'];snapshot['projects'].append(pid)
        originals={item['sha256']:item for item in selected['files']}
        for doc in incoming['documents']:
            require(doc.get('extraction_status')=='EXTRACTED','SOURCE_EXTRACTION_NOT_COMPLETE')
            original=api.request('GET',f"/api/documents/{doc['id']}/original").content
            require(sha256(original).hexdigest()==doc['sha256'],'UPLOADED_ORIGINAL_MISMATCH')
            require(original==read_original(originals[doc['sha256']]),'LOCAL_ORIGINAL_MISMATCH')
            knowledge=api.json('GET',f"/api/documents/{doc['id']}/knowledge")
            require(bool(knowledge.get('markdown')) and bool(knowledge.get('blocks')) and bool(knowledge.get('sidecars')),'KNOWLEDGE_MISSING')
            snapshot['documents'].append({'id':doc['id'],'sha256':doc['sha256'],'sidecars':knowledge['sidecars']})
        check(report,'fresh_upload_originals_and_knowledge',project=selected['code'],fresh_hashes=incoming['fresh_hash_count'])
        analysis=api.json('POST',f'/api/projects/{pid}/analyze',json={'mode':'rules'})
        require(analysis['ai_used'] is False,'RULES_MISLABELED_AS_AI')
        require(any(row.get('current_refs') and row.get('historical_refs') for row in analysis['rows']),'HISTORICAL_CONNECTION_MISSING')
        if selected['code']=='N01':
            notice=next(row for row in analysis['rows'] if row['id']=='notice')
            require('14 calendar days' in notice['current'] and '7 calendar days' not in notice['current'],'APPROVED_AMENDMENT_PRECEDENCE_FAILED')
            require(any('Rev02' in str(ref.get('revision')) for ref in notice.get('excluded_refs',[])),'DRAFT_AMENDMENT_NOT_EXCLUDED')
        graph=api.json('GET',f'/api/projects/{pid}/knowledge')
        require(any(edge.get('source_ref') for edge in graph['edges']),'GRAPH_SOURCE_REF_MISSING')
        check(report,'rules_comparison_and_source_graph',project=selected['code'],rows=len(analysis['rows']))
        # Re-upload identical bytes before approval: preserve source identity/dedup.
        first=selected['files'][0];before=len(api.json('GET',f'/api/projects/{pid}/documents'))
        duplicate=api.json('POST',f'/api/projects/{pid}/upload',files={'files':(first['name'],read_original(first),'application/octet-stream')})['documents'][0]
        require(duplicate.get('duplicate') is True and len(api.json('GET',f'/api/projects/{pid}/documents'))==before,'DUPLICATE_INPUT_NOT_IDEMPOTENT')
        for kind in ['itb','risk','slides']:
            draft=api.json('POST',f'/api/projects/{pid}/drafts',json={'kind':kind,'mode':'rules'})
            api.request('POST',f"/api/drafts/{draft['id']}/export",expected=(409,))
            rows=deepcopy(draft['rows'])
            if kind=='slides':rows[0]['title']='프로젝트 개요 검토'
            else:rows[0]['rationale']='시험 계정 검증: 원문 근거를 확인한 편집 기록'
            draft=api.json('PATCH',f"/api/drafts/{draft['id']}",json={'rows':rows,'revision':draft['revision']})
            draft=api.json('POST',f"/api/drafts/{draft['id']}/approve",json={'confirmed':True,'revision':draft['revision']})
            data=api.request('POST',f"/api/drafts/{draft['id']}/export").content
            signature=office_signature(data,kind)
            extension='pptx' if kind=='slides' else 'xlsx'
            path=output_dir/f"{selected['code']}-{kind}.{extension}";path.write_bytes(data)
            snapshot['drafts'].append({'id':draft['id'],'kind':kind,'revision':draft['revision'],
                                      'approval_id':draft['approval_id'],'input_fingerprint':draft['input_fingerprint'],
                                      'file_sha256':sha256(data).hexdigest(),'path':str(path),'signature':signature})
            check(report,'review_approval_and_actual_output',project=selected['code'],kind=kind,**signature)
    # Rehash the exact selected bytes after processing; never modify input originals.
    for selected in past+new:
        for item in selected['files']:read_original(item)
    check(report,'selected_original_hashes_unchanged')
    return snapshot


def capture_output_objects(store, snapshot):
    events=[row['payload'] for row in store.list('event')]
    snapshot['stored_outputs']=[]
    for draft in snapshot['drafts']:
        matches=[event for event in events if event.get('action')=='output_generated'
                 and event.get('draft_id')==draft['id'] and event.get('output_sha256')==draft['file_sha256']]
        require(bool(matches),'STORED_OUTPUT_EVENT_MISSING')
        event=matches[0]
        data=store.download_bytes(event['storage_path'])
        require(sha256(data).hexdigest()==draft['file_sha256'],'STORED_OUTPUT_HASH_MISMATCH')
        snapshot['stored_outputs'].append({'path':event['storage_path'],'sha256':draft['file_sha256']})


def recheck_persistence(api, snapshot: dict, report: dict, store):
    require(len(snapshot.get('stored_outputs',[]))==len(snapshot['drafts'])>0,'STORED_OUTPUT_CHECKPOINT_MISSING')
    # Read exact pre-restart objects before export can regenerate missing files.
    for output in snapshot['stored_outputs']:
        try:data=store.download_bytes(output['path'])
        except Exception:raise DemoError('RESTART_STORED_OUTPUT_UNAVAILABLE') from None
        require(sha256(data).hexdigest()==output['sha256'],'RESTART_STORED_OUTPUT_HASH_MISMATCH')
    check(report,'exact_output_objects_persisted_before_regeneration',outputs=len(snapshot['stored_outputs']))
    for pid in snapshot['projects']:api.json('GET',f'/api/projects/{pid}')
    for doc in snapshot['documents']+snapshot['historical_documents']:
        data=api.request('GET',f"/api/documents/{doc['id']}/original").content
        require(sha256(data).hexdigest()==doc['sha256'],'RESTART_ORIGINAL_HASH_MISMATCH')
        knowledge=api.json('GET',f"/api/documents/{doc['id']}/knowledge")
        require(knowledge['sidecars']==doc['sidecars'] and knowledge.get('markdown'),'RESTART_KNOWLEDGE_MISSING')
    for saved in snapshot['drafts']:
        draft=api.json('GET',f"/api/drafts/{saved['id']}")
        require(draft['status']=='approved' and all(draft[key]==saved[key] for key in ['revision','approval_id','input_fingerprint']),'RESTART_APPROVAL_CHANGED')
        data=api.request('POST',f"/api/drafts/{saved['id']}/export").content
        require(office_signature(data,saved['kind'])==saved['signature'],'RESTART_OUTPUT_CONTENT_CHANGED')
    check(report,'persisted_approved_outputs',documents=len(snapshot['documents']),outputs=len(snapshot['drafts']))


def verify_sidecars(api, settings, owner, snapshot, report):
    store=SupabaseStore(settings.supabase_url,settings.publishable_key,api.token)
    try:
        for doc in snapshot['documents']+snapshot['historical_documents']:
            md=store.download_bytes(doc['sidecars']['markdown']).decode('utf-8')
            sidecar=json.loads(store.download_bytes(doc['sidecars']['json']))
            require(md and sidecar['sha256']==doc['sha256'] and sidecar['id']==doc['id'],'SIDECAR_STORAGE_MISMATCH')
    finally:store.close()
    check(report,'actual_user_jwt_sidecar_downloads')


def list_owned_storage(http, url, headers, owner):
    require(str(UUID(owner))==owner,'CLEANUP_OWNER_INVALID')
    pending=[owner];visited=set();objects=set()
    while pending:
        prefix=pending.pop()
        require(prefix not in visited and len(visited)<10000,'CLEANUP_DIRECTORY_LIMIT')
        visited.add(prefix);seen=set()
        for offset in range(0,100000,100):
            response=http.request('POST',f'{url}/storage/v1/object/list/rebuild-agent',headers=headers,
                json={'prefix':prefix,'limit':100,'offset':offset,'sortBy':{'column':'name','order':'asc'}},timeout=20)
            require(response.status_code==200,'CLEANUP_STORAGE_LIST_UNAVAILABLE')
            page=response.json();require(isinstance(page,list) and len(page)<=100,'CLEANUP_STORAGE_LIST_INVALID')
            for item in page:
                name=item.get('name')
                require(isinstance(name,str) and name not in {'','.','..'} and not any(c in name for c in '/\\\x00')
                        and name not in seen,'CLEANUP_STORAGE_PATH_INVALID')
                seen.add(name);path=prefix+'/'+name
                require(path.startswith(owner+'/'),'CLEANUP_STORAGE_OWNER_MISMATCH')
                if item.get('id') is None and item.get('metadata') is None:pending.append(path)
                else:objects.add(path)
            if len(page)<100:break
        else:raise DemoError('CLEANUP_STORAGE_LIST_LIMIT')
    return sorted(objects)


def cleanup_created_users(http, url, secret_key, users, run_id='') -> bool:
    # Preserve recovery metadata if any physical Storage inventory/deletion fails.
    if not run_id:return False
    headers=_service_headers(secret_key);complete=True
    for owner in users:
        try:
            response=http.request('GET',f'{url}/auth/v1/admin/users/{owner}',headers=headers,timeout=20)
            require(response.status_code==200,'CLEANUP_OWNER_UNAVAILABLE')
            metadata=response.json().get('user_metadata',{})
            require(metadata.get('run_id')==run_id and metadata.get('created_by')=='rebuild-agent-live-probe','CLEANUP_OWNER_MISMATCH')
            # Physical paths can differ from row paths; failed ingestion may have no row.
            paths=list_owned_storage(http,url,headers,owner)
            for offset in range(0,len(paths),100):
                response=http.request('DELETE',f'{url}/storage/v1/object/rebuild-agent',headers=headers,
                    json={'prefixes':paths[offset:offset+100]},timeout=20)
                require(response.status_code in (200,204),'CLEANUP_STORAGE_FAILED')
            require(not list_owned_storage(http,url,headers,owner),'CLEANUP_STORAGE_NOT_EMPTY')
            response=http.request('DELETE',f'{url}/rest/v1/rb_entities',headers=headers,params={'owner_id':f'eq.{owner}'},timeout=20)
            require(response.status_code in (200,204),'CLEANUP_ROWS_FAILED')
            response=http.request('DELETE',f'{url}/auth/v1/admin/users/{owner}',headers=headers,timeout=20)
            require(response.status_code in (200,204),'CLEANUP_USER_FAILED')
        except Exception:complete=False
    return complete


class OwnedServer:
    def __init__(self, run_id, port, readiness_timeout=60):
        if type(readiness_timeout) not in (int,float) or not 1 <= readiness_timeout <= 240:
            raise ValueError("INVALID_SERVER_READINESS_BOUND")
        self.readiness_timeout=readiness_timeout
        self.last_start_seconds=None
        self.run_id,self.port,self.process,self.log=run_id,port,None,None
        self.url=f'http://127.0.0.1:{port}'
    def start(self):
        with socket.socket() as probe:probe.bind(('127.0.0.1',self.port))
        logs=ROOT/'ops/private';logs.mkdir(parents=True,exist_ok=True)
        self.log=(logs/f'demo-live-{self.run_id}.log').open('ab')
        env={**os.environ,'REBUILD_LIVE_VERIFY_RUN_ID':self.run_id,'REBUILD_ENV':'local'}
        started=time.monotonic()
        self.process=subprocess.Popen([sys.executable,'-m','uvicorn','ops.demo_live_server:app','--host','127.0.0.1','--port',str(self.port),'--no-access-log'],
            cwd=ROOT,env=env,stdout=self.log,stderr=self.log,creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
        deadline=started+self.readiness_timeout
        with httpx.Client(timeout=2,trust_env=False) as http:
            while time.monotonic()<deadline:
                if self.process.poll() is not None:raise DemoError('OWNED_SERVER_EXITED')
                try:
                    response=http.get(self.url+'/__rebuild_live_verify_identity')
                    if response.status_code==200 and response.json()=={'run_id':self.run_id,'boundary':'REAL_SUPABASE_USER_JWT','product':'RE:Build Agent'}:
                        self.last_start_seconds=round(time.monotonic()-started,2)
                        return self.process.pid
                except (httpx.RequestError,ValueError):pass
                time.sleep(.2)
        self.last_start_seconds=round(time.monotonic()-started,2)
        raise DemoError('OWNED_SERVER_NOT_READY')
    def stop(self):
        if self.process and self.process.poll() is None:
            self.process.terminate()
            try:self.process.wait(timeout=15)
            except subprocess.TimeoutExpired:self.process.kill();self.process.wait(timeout=5)
        if self.log:self.log.close()


def main() -> int:
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--past',nargs='+',default=['P01','P06'])
    parser.add_argument('--new',nargs='+',choices=['N01','N02'],default=['N01','N02'])
    parser.add_argument('--with-ai',action='store_true')
    parser.add_argument('--port',type=int,default=8792)
    args=parser.parse_args()
    report={'product':'RE:Build Agent','boundary':'REAL_SUPABASE_USER_JWT','at':datetime.now(timezone.utc).isoformat(),
            'status':'PENDING','checks':[],'cleanup_complete':True,'product_complete':False,'ai_verified':False}
    users=[];server=None;run_id=uuid4().hex;stage='configuration'
    report['run_id']=run_id
    load_environment();settings=Settings.from_environment();secret_key=os.environ.get('SUPABASE_SECRET_KEY','')
    with httpx.Client(timeout=20,trust_env=False) as admin:
        try:
            settings.validate();require(bool(secret_key),'SUPABASE_SECRET_KEY_REQUIRED')
            stage='schema_and_rls_preflight'
            report['preflight']=verify_live(settings.supabase_url,settings.publishable_key,secret_key)
            require(report['preflight']['status']=='PASSED','SUPABASE_PREFLIGHT_NOT_PASSED')
            entries=catalog();past=select_projects(entries,args.past,'historical');new=select_projects(entries,args.new,'current')
            stage='create_isolated_users'
            credentials=[]
            for index in range(2):
                email=f'rebuild-dataset-{run_id}-{index}@example.com';password=secrets.token_urlsafe(30)
                owner=_create_test_user(admin,settings.supabase_url,secret_key,email,password,run_id,20)
                users.append(owner);credentials.append((email,password))
            stage='owned_server_start';server=OwnedServer(run_id,args.port);first_pid=server.start()
            output_dir=ROOT/'ops/runtime/demo'/run_id
            with httpx.Client(base_url=server.url,timeout=180,follow_redirects=False,trust_env=False) as client:
                api=ProductAPI(client);api.login(*credentials[0])
                stage='actual_input_flow';snapshot=run_test_flow(api,past,new,output_dir,report)
                verify_sidecars(api,settings,users[0],snapshot,report)
                store=SupabaseStore(settings.supabase_url,settings.publishable_key,api.token)
                try:capture_output_objects(store,snapshot)
                finally:store.close()
                report['snapshot']=snapshot
                stage='second_user_isolation';other=ProductAPI(client);other.login(*credentials[1])
                require(other.json('GET','/api/projects')==[],'CROSS_USER_PROJECT_LEAK')
                other.request('GET',f"/api/documents/{snapshot['documents'][0]['id']}/original",expected=(404,))
                other.request('POST',f"/api/drafts/{snapshot['drafts'][0]['id']}/export",expected=(404,))
                other.logout();check(report,'product_cross_user_access_denied')
                if args.with_ai:
                    stage='local_oauth_product_call'
                    answer=api.json('POST',f"/api/projects/{snapshot['projects'][0]}/analyze",json={'mode':'ai','question':'승인된 통지 기한과 과거 사례의 차이를 원문 근거로 설명하세요.'})
                    execution=answer.get('ai_insight',{}).get('execution',{})
                    require(answer.get('ai_used') and execution.get('source')=='MODEL_CALL'
                            and execution.get('auth_mode')=='CHATGPT_OAUTH','ACTUAL_OAUTH_MODEL_CALL_NOT_PROVED')
                    report['ai_verified']=True;report['ai_execution']=execution;report['ai_usage']=answer['ai_insight'].get('usage')
                api.logout()
            stage='actual_process_restart';server.stop();second_pid=server.start()
            require(second_pid!=first_pid,'PROCESS_NOT_RESTARTED')
            with httpx.Client(base_url=server.url,timeout=180,trust_env=False) as client:
                api=ProductAPI(client);api.login(*credentials[0])
                store=SupabaseStore(settings.supabase_url,settings.publishable_key,api.token)
                try:recheck_persistence(api,snapshot,report,store)
                finally:store.close()
                verify_sidecars(api,settings,users[0],snapshot,report);api.logout()
            check(report,'actual_process_restart_and_new_login',first_pid=first_pid,second_pid=second_pid)
            report['status']='PASSED_WITH_LOCAL_AI' if report['ai_verified'] else 'PASSED_RULES_FLOW_AI_NOT_TESTED'
        except (Exception,KeyboardInterrupt) as error:
            report.update(status='BLOCKED' if stage=='schema_and_rls_preflight' else 'FAILED',stage=stage,
                          error_code=str(error) if isinstance(error,DemoError) else 'SAFE_DIAGNOSTIC_WITHHELD')
        finally:
            if server:server.stop()
            if users:
                report['cleanup_complete']=cleanup_created_users(admin,settings.supabase_url,secret_key,users,run_id)
                if not report['cleanup_complete']:report['status']='CLEANUP_REQUIRED';report['cleanup_owner_ids']=users
    path=write_report(report,'live-verification')
    print(json.dumps({'status':report['status'],'stage':report.get('stage'),'report':str(path.relative_to(ROOT)),
                      'cleanup_complete':report['cleanup_complete'],'ai_verified':report['ai_verified']},ensure_ascii=False))
    return 0 if report['status'].startswith('PASSED') and report['cleanup_complete'] else 1


if __name__=='__main__':raise SystemExit(main())
